# -*- coding: utf-8 -*-
"""组装 Word 报告:深度分析 + 简洁排版。
原则:去低价值噪音,但保留透彻分析;加粗仅用于结论性短语、不给普通数字上色;
图表只选有信息量的;直接用有意义口径的数据。
数据源:stats.json / stats2.json / synth.json / validation.json / cleaning_log.json。
"""
from __future__ import annotations
import json
from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = Path("analysis/report/outputs"); CH = OUT / "charts"
S = json.load(open(OUT / "stats.json", encoding="utf-8"))
S2 = json.load(open(OUT / "stats2.json", encoding="utf-8"))
SY = json.load(open(OUT / "synth.json", encoding="utf-8"))
VAL = json.load(open(OUT / "validation.json", encoding="utf-8"))
CLEAN = json.load(open(OUT / "cleaning_log.json", encoding="utf-8"))

CJK = "微软雅黑"
NAVY = RGBColor(0x2F, 0x55, 0x97); GREY = RGBColor(0x59, 0x59, 0x59)
GREEN = "548235"; AMBER = "BF8F00"; RED = "C00000"
HDR_FILL = "2F5597"; ALT_FILL = "F2F5FA"

def brl(x):
    if x is None: return "/"
    x = float(x)
    if abs(x) >= 1e8: return f"{x/1e8:,.1f}亿"
    if abs(x) >= 1e4: return f"{x/1e4:,.0f}万"
    return f"{x:,.0f}"
def num(x): return "/" if x is None else f"{x:,.0f}"

def cx(s):
    """去破折号:中文破折号(——/—)改为逗号,en-dash 改连字符,清理重复逗号。"""
    if not isinstance(s, str): return s
    s = s.replace("——", ",").replace("—", ",").replace("–", "-")
    while ",," in s: s = s.replace(",,", ",")
    return s.replace(",。", "。").replace(",、", "、").replace(",,", ",")

def _font(run, size=None, bold=None, color=None, name=CJK):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), name)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = color

def _styles(doc):
    st = doc.styles["Normal"]; st.font.name = CJK; st.font.size = Pt(10.5)
    rf = st.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), CJK)
    for i, sz in [(1, 17), (2, 13)]:
        h = doc.styles[f"Heading {i}"]; h.font.name = CJK; h.font.size = Pt(sz)
        h.font.bold = True; h.font.color.rgb = NAVY
        rf = h.element.get_or_add_rPr().get_or_add_rFonts()
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), CJK)

def para(doc, text="", size=10.5, color=None, align=None, after=6, before=0, italic=False, indent=None):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if indent: p.paragraph_format.left_indent = Inches(indent)
    if text:
        r = p.add_run(cx(text)); _font(r, size, None, color); r.font.italic = italic
    return p

def body(doc, text, size=10.5):
    """分析正文:整段普通字重(数字不加粗),按换行分段。"""
    for seg in str(text).split("\n"):
        seg = seg.strip()
        if seg: para(doc, seg, size=size, after=7)

def heading(doc, text, level=1):
    h = doc.add_heading(level=level); r = h.add_run(cx(text)); _font(r, None, True, NAVY); return h

def keypoint(doc, text, size=10.5):
    """长分析前的加粗要点结论(便于快速理解)。"""
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.space_before = Pt(2)
    r0 = p.add_run("核心结论  "); _font(r0, size, bold=True, color=NAVY)
    r = p.add_run(cx(text)); _font(r, size, bold=True)
    return p

def takeaway(doc, lead, rest="", size=10.5):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(cx(lead)); _font(r, size, bold=True)
    if rest: r2 = p.add_run(cx(rest)); _font(r2, size)
    return p

def add_image(doc, name, width=5.9, caption=None):
    fp = CH / name
    if not fp.exists(): para(doc, f"[缺图 {name}]", color=GREY, italic=True); return
    doc.add_picture(str(fp), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption: para(doc, caption, size=8.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

def _shade(cell, fill):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(sh)

def table(doc, headers, rows, widths, font=9, align_right=None, cell_colors=None):
    align_right = align_right or set()
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ""; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h)); _font(r, font, True, RGBColor(0xFF, 0xFF, 0xFF)); _shade(c, HDR_FILL)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""; p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j in align_right else WD_ALIGN_PARAGRAPH.LEFT
            col = RGBColor.from_string(cell_colors[(i, j)]) if (cell_colors and (i, j) in cell_colors) else None
            r = p.add_run(cx(str(val))); _font(r, font, bold=(col is not None), color=col)
            if i % 2 == 1: _shade(cells[j], ALT_FILL)
    for j, w in enumerate(widths):
        for row in t.rows: row.cells[j].width = Inches(w)
    return t

def add_toc(doc):
    p = doc.add_paragraph(); run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    tt = OxmlElement("w:t"); tt.text = "【在 Word 中右键 → 更新域 生成目录】"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, instr, f2, tt, f3): run._r.append(e)
    _font(run, 10.5, color=GREY)

def page_break(doc): doc.add_page_break()

# ════════════════════════════════════════════════════════════
doc = Document(); sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"): setattr(sec, m, Inches(1))
_styles(doc)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("巴西政府采购市场分析  ·  "); _font(r, 8, color=GREY)
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); fp._p.append(fld)

mo = S["market_overview"]; tot = mo["totals"]; op = S["opportunities"]
cp = S["competition"]; conc = cp["concentration"]; fn = S2["funnel"]
loc = S2["localization"]; fd = loc["foreign_direct"]; lm = loc["localized_mnc_heuristic"]
mh = S2["match"]; dm = S["data_mining"]
nar = SY["section_narratives"]
open_eff = next(s["数量"] for s in fn["stages"] if s["阶段"].startswith("③"))
open_prio = next(s["数量"] for s in fn["stages"] if s["阶段"].startswith("⑤"))
dim = {r["主维度"]: r for r in mo["by_dimension"]}
odim = {r["主维度"]: r for r in op["open_by_dimension"]}

# ── 封面 ──
for _ in range(4): para(doc, "")
para(doc, "巴西政府采购(PNCP)市场分析", size=26, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
doc.paragraphs[-1].runs[0].font.bold = True
para(doc, "面向中国企业的商机、市场与竞品研判", size=13, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=36)
win = S["meta"].get("采集窗口_发布", ["", ""])
para(doc, f"数据来源:巴西全国政府采购门户 PNCP   |   采集窗口:{str(win[0])[:10]} 至 {str(win[1])[:10]}",
     size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para(doc, f"样本:招标 {S['meta']['n_招标']:,} · 采购明细 {S['meta']['n_明细']:,} · 合同 {S['meta']['n_合同']:,} · 年度计划 {S['meta']['n_PCA']:,} · 机构 {S['meta']['n_机构']:,}",
     size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para(doc, "方法:数据清洗 → 统计分析 → 数据挖掘 → 多视角洞察(对抗式数字校验)   |   报告日期:2026-06-08",
     size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
page_break(doc)
heading(doc, "目录", 1); add_toc(doc); page_break(doc)

# ════ 一、执行摘要 ════
heading(doc, "一、执行摘要", 1)
para(doc, "巴西政府采购是一个高频、分散、由本地法人主导的市场。机会真实且充裕,但取得本地身份是中标的前提;"
          "中企应以「本地化主体 + 电子标跑量 + 高端制造/医疗/数字经济选择性突破 + 用年度计划提前卡位」为主线。", after=8)
EXEC = [
    ("高频分散的中小标市场。", f"一周 8,357 条招标、披露 {brl(tot['总额_BRL'])} BRL,金额中位仅 {brl(tot['中位数_BRL'])}、均值 {brl(tot['均值_BRL'])}——多数为中小标,少数大额标拉高均值。"),
    ("可投机会充裕且滚动。", f"在投且属六大维度的有效机会 {open_eff:,} 条,收窄到电子标/已披露后的优先清单 {open_prio:,} 条;机会按周滚动,无需苦等大项目。"),
    ("本地化是中标的分水岭。", f"外企以纯境外身份直接中标仅 {fd['n_contracts']} 笔(0.04%);完成本地化的外资子公司至少 {lm['n_suppliers']} 家、{lm['n_contracts']} 笔中标(含徐工、同方威视),本地法人占约 99% 金额。"),
    ("未来管道清晰可提前布局。", f"2026 年度采购计划 {op['pca_by_year'][2]['n']:,} 项、约 {brl(op['pca_by_year'][2]['total'])} BRL,12 月为放量高峰。"),
    ("决策与资金高度下沉。", "市级机构约占 76% 笔数、62% 金额,渠道必须落到市、州两级。"),
    ("竞争呈长尾而非垄断。", f"{conc['n_unique_suppliers']:,} 家供应商,基尼≈0.97 但 HHI 仅 {conc['HHI']:.0f}——头部金额由一次性大单驱动,新进入者无需正面挑战在位巨头。"),
]
for lead, rest in EXEC: takeaway(doc, lead, rest)
page_break(doc)

# ════ 二、市场全景 ════
heading(doc, "二、市场全景", 1)
keypoint(doc, "市场高频分散、以中小标为主;价值集中在医疗医药与大宗商贸,采购决策权高度下沉到市级。")
body(doc, nar["market"])
para(doc, "")
add_image(doc, "01_dimension.png", caption="图1 六大维度:招标数(蓝柱·左轴)与披露金额(橙柱·右轴)")
para(doc, f"图1 用蓝柱(招标数,左轴)与橙柱(披露金额,右轴)并排,直观呈现「量价分离」:"
          f"大宗商贸蓝柱最高({num(dim['大宗商贸']['n'])} 条)但橙柱中等,说明数量多、单标偏小(同质化、拼价格);"
          f"医疗医药蓝柱不高但橙柱很高、金额中位达 {brl(dim['医疗医药']['median'])},是专业化、毛利空间更大的赛道;"
          f"数字经济仅 {num(dim['数字经济']['n'])} 条却由少数高价标支撑;跨境电商样本极小。对中企的含义:做大宗商贸要拼成本规模,"
          "做医疗/数字经济则拼产品与方案、单笔回报更高。", size=10, after=6)
rows = [[r["主维度"], num(r["n"]), brl(r["total"]), brl(r["median"])] for r in mo["by_dimension"]]
table(doc, ["业务维度", "招标数", "披露金额", "金额中位"], rows, widths=[1.7, 1.2, 1.6, 1.4], align_right={1, 2, 3})
rows = [[r["大区"], num(r["n"]), brl(r["total"]), brl(r["median"])] for r in mo["by_region"]]
table(doc, ["大区", "招标数", "披露金额", "金额中位"], rows, widths=[1.6, 1.2, 1.5, 1.4], align_right={1, 2, 3})
takeaway(doc, "东南区做基本盘、北部/中西部抓大单。", "东南区招标数最多、靠近圣保罗工业物流核心;北部笔数少但中位金额最高(被个别大额医疗认证标拉动)。")
takeaway(doc, "采购形式对外企友好。", "电子拍卖(Pregão Eletrônico)为绝对主力,全程线上、按价竞标,降低了无本地人脉者的参与门槛。")
page_break(doc)

# ════ 三、中企商机 ════
heading(doc, "三、中企商机", 1)
keypoint(doc, "可投机会充裕且按周滚动;主攻电子拍卖与 SRP,以设备/产品类切入数字经济、医疗医药、高端制造,并用年度计划提前卡位。")
body(doc, nar["opportunity"])
para(doc, "")
heading(doc, "3.1 可投机会聚焦", 2)
para(doc, f"入库的 8,357 条已通过 PNCP 相关性过滤;剔除已过期与未分类后,在投有效机会 {open_eff:,} 条,"
          f"再叠加已披露预算与电子/SRP 友好方式,优先可投清单 {open_prio:,} 条——这是中企应集中投入的高质量靶面。", after=6)
add_image(doc, "14_funnel.png", caption="图2 可投机会逐层聚焦")
takeaway(doc, "主攻电子标两条线。", f"电子拍卖在投 {op['friendly_modes']['电子拍卖_Pregao_Eletronico_在投']:,} 条用于跑量练兵;价格登记 SRP 在投 {op['friendly_modes']['价格登记SRP_在投']:,} 条是框架协议,一次入围、长期按需供货,可摊薄投标成本。")
takeaway(doc, "数字经济少而精。", f"在投仅 {odim['数字经济']['n']} 条却达 {brl(odim['数字经济']['total'])} BRL,适合以安防监控、政务/企业级 IT 设备、系统集成差异化切入,绕开本地低价红海。")
takeaway(doc, "卖产品而非拿工程。", f"医疗医药(在投 {odim['医疗医药']['n']} 条/{brl(odim['医疗医药']['total'])})与高端制造(车辆、工程机械、电力设备)以供货为主,门槛低于需本地施工资质的工程类(Obras)。")
para(doc, "")
heading(doc, "3.2 高价值在投机会(经中企可承接性筛选)", 2)
para(doc, "筛选口径四步:① 时效为「还能投/临近截止」(在可投标窗口内);② 已披露预估金额;③ 按金额降序;"
          "④ 关键一步,按「中企可承接性」过滤,只保留中国企业能供货/承接、与我方主营可对接的标。"
          "纯本地工程与服务(土建施工、道路养护、人员运输、保洁安保、文体活动、金融服务等)即便金额再高,中企也无从承接,故不计入。", size=10, after=6)
HVA = json.load(open(OUT / "highvalue_assessed.json", encoding="utf-8"))
HVR = {r["PNCP编号"]: r for r in json.load(open(OUT / "highvalue_raw.json", encoding="utf-8"))}
def feasibility(r):
    if r.get("长期机会") == "是": return "长期挂网(随到随审)"
    dd = r.get("剩余天数")
    if dd is None: return "未知"
    if dd >= 180: return "充裕(可先本地化再投)"
    if dd >= 90:  return "偏紧(需已本地化/快速合资)"
    return "紧张(限已有本地法人/伙伴)"
items = []
for a in HVA:
    r = HVR.get(a["pncp_id"])
    if r: items.append({**a, **r, "时效可行性": feasibility(r)})
n_hi = sum(1 for x in items if x["可承接性"] == "高")
n_md = sum(1 for x in items if x["可承接性"] == "中")
n_lo = sum(1 for x in items if x["可承接性"] == "低")
match = sorted([x for x in items if x["可承接性"] in ("高", "中")], key=lambda x: -(x.get("valor_brl_clean") or 0))
para(doc, f"在金额最高的 {len(items)} 条在投标中,经判定可由中企承接的有 {n_hi + n_md} 条(可承接 {n_hi}、部分可承接 {n_md}),"
          f"其余 {n_lo} 条为纯本地工程/服务、与我方无关。下表为可承接的高价值标(按金额,标的为概括非截取):", size=10, after=4)
SUPC2 = {"高": GREEN, "中": AMBER}
rows, colors = [], {}
for i, x in enumerate(match[:14]):
    rows.append([str(x.get("梗概") or "")[:24], brl(x.get("valor_brl_clean")), str(x.get("州") or ""),
                 str(x.get("剩余天数") or ""), x["时效可行性"], x["可承接性"]])
    colors[(i, 5)] = SUPC2.get(x["可承接性"], "595959")
table(doc, ["标的梗概(概括)", "金额", "州", "剩余天数", "时效可行性(对接本地化)", "可承接"], rows,
      widths=[2.05, 0.95, 0.45, 0.65, 1.6, 0.6], font=8, align_right={1, 3}, cell_colors=colors)
para(doc, "时效可行性结合本地化周期判断:本地化通常需 3 至 6 个月,故剩余天数不足 90 天的标只适合「已有本地法人或本地伙伴」的中企直接参与;"
          "未本地化的中企面对这类短窗口标,应走「与本地企业组联合体」,或转向年度计划(PCA)与长期挂网(Credenciamento/SRP)从容布局。", size=9.5, color=GREY, after=8)
heading(doc, "3.3 未来管道:年度采购计划(PCA)", 2)
para(doc, "PCA 是政府提前公示「买什么、何时买」的年度意向,等于一份招标日历。中企可据此在正式公告前数月接触机构、备样品认证与本地仓储。", after=6)
add_image(doc, "08_pca_pipeline.png", width=6.1, caption="图3 年度采购计划:左为按年度规模、右为按类别 Top6")
pcat = {r["物/服"]: r["total"] for r in op["pca_top_categories"]}
para(doc, f"左图(按年度)看规模与时点:2026 年一年就有 {op['pca_by_year'][2]['n']:,} 项、约 {brl(op['pca_by_year'][2]['total'])} BRL,"
          f"远高于 2025 年的 {brl(op['pca_by_year'][1]['total'])};这是因为各机构主要为当年/次年报计划,2027 年仅 {op['pca_by_year'][3]['n']:,} 项、仍在陆续补登,"
          "可当作「领先指标」持续跟踪。结论是:绝大多数未来商机都压在 2026 年,现在正是介入窗口。", size=10, after=6)
para(doc, f"右图(按类别)看可触及性:服务类金额最大(约 {brl(pcat.get('Serviço'))}),但其中很大一部分是难以贸易的本地人力服务;"
          f"对中企真正可切入的是物资/材料(约 {brl(pcat.get('Material'))})与信息通信 TIC 方案(约 {brl(pcat.get('Soluções de TIC'))}),"
          f"以及其中的医疗设备/手术器械、计算机、车辆等实物品类;工程及工程服务(约 {brl(pcat.get('Obras e Serviços de Engenharia'))})需本地施工资质,优先级较低。"
          "也就是说,金额最大的类别不等于机会最大的类别,应按「可贸易、可出口」筛选品类。", size=10, after=6)
rows = [[str(r["计划年度"]), num(r["n"]), brl(r["total"])] for r in op["pca_by_year"]]
table(doc, ["计划年度", "计划项数", "预估金额"], rows, widths=[1.5, 1.7, 1.8], align_right={1, 2})
takeaway(doc, "用 PCA 提前 3 至 6 个月卡位。", "锁定 2026 年的 TIC、医疗设备、计算机、车辆等实物品类,按目标机构与期望采购月接触备货,赶在 12 月放量高峰前完成资质。")
page_break(doc)

# ════ 四、竞品情报 ════
heading(doc, "四、竞品情报", 1)
keypoint(doc, "市场呈长尾而非垄断,纯境外身份几乎无缘中标;注册本地法人(PJ)、组建本地子公司是中标的硬门槛。")
body(doc, nar["competition"])
para(doc, "")
heading(doc, "4.1 供应商集中度", 2)
add_image(doc, "09_supplier_lorenz.png", width=4.6, caption="图4 中标金额集中度(洛伦兹曲线)")
para(doc, f"基尼系数≈0.97 说明金额分配极不均(大量供应商只中得小额标),但 HHI 仅约 {conc['HHI']:.0f}(中等)说明并非单一巨头垄断——"
          f"Top10 占金额 {conc['top10_value_share_pct']:.0f}% 的份额,主要由若干一次性巨额合同(州级公路联合体、车队管理、银行服务等)瓜分。"
          "对中企的含义是:这是一个「长尾极散、头部由项目制大单驱动」的市场,新进入者可以从中小标切入积累信用,而不必正面挑战在位者。", size=10, after=8)
rows = [[str(r["中标供应商"])[:32], brl(r["total"]), num(r["n"])] for r in cp["top_suppliers_by_value"][:10]]
table(doc, ["头部中标供应商", "合同金额", "合同数"], rows, widths=[3.6, 1.4, 1.0], font=8.5, align_right={1, 2})
para(doc, "头部供应商多数 n=1,印证「项目制大单」而非长期复购垄断。", size=9.5, color=GREY, after=8)
heading(doc, "4.2 本地化:中标的分水岭", 2)
add_image(doc, "15_localization.png", width=6.1, caption="图5 外企直接 vs 本地化外资子公司;中标金额按主体")
para(doc, f"「本地化是变现门槛」可量化验证:以纯境外主体直接投标,全样本仅 {fd['n_contracts']} 笔中标、{fd['n_suppliers']} 家企业、{brl(fd['value_BRL'])} BRL,"
          f"且几乎全是学术数据库、科研试剂等无法本地化生产的特例。而完成本地化(注册巴西法人)的外资跨国企业,按名称启发式至少可识别 "
          f"{lm['n_suppliers']} 家、拿下 {lm['n_contracts']} 笔合同、合计 {brl(lm['value_BRL'])} BRL——中标家数与笔数比「裸投标」高出一个数量级。"
          f"从主体看,中标金额约 99% 归本地法人(PJ),自然人与其他主体合计不足 1%。", size=10, after=6)
para(doc, "已识别的本地化外资子公司中包含中国企业——徐工(XCMG BRASIL)、同方威视(NUCTECH DO BRASIL)等已在巴西中标,"
          "证明「中企本地化 → 中标」已有先例可循。部分样例:", size=10, after=4)
rows = [[str(r.get("中标供应商") or "")[:36], str(r.get("pais") or "BRA"), brl(r.get("valor_brl_clean")), str(r.get("州") or "")]
        for r in loc["examples_localized_mnc"][:10]]
table(doc, ["本地化外资子公司(中标企业)", "登记国", "合同金额", "州"], rows, widths=[3.4, 0.9, 1.4, 0.6], font=8, align_right={2})
para(doc, "注:本地化子公司为名称启发式识别(下限估计),真实数量更多。", size=8.5, color=GREY, after=8)
heading(doc, "4.3 可追溯案例:招标 → 中标企业 → 金额", 2)
para(doc, f"按 linked_contratacao_pncp_id 关联,本快照中 {mh['n_matched']} 笔合同可匹配到对应招标,还原完整成交链路"
          "(匹配率偏低是因为合同多为历史授标、与本周招标时间窗不同)。样例 Top 8:", size=10, after=4)
rows = [[str(r.get("主维度") or ""), str(r.get("标的(中文梗概)") or "")[:20], str(r.get("中标供应商") or "")[:22],
         brl(r.get("合同金额BRL")), str(r.get("州") or "")] for r in S2["match_records"][:8]]
table(doc, ["维度", "招标标的", "中标企业", "合同金额", "州"], rows, widths=[0.85, 1.8, 1.85, 1.1, 0.5], font=8, align_right={3})
page_break(doc)

# ════ 五、数据挖掘发现 ════
heading(doc, "五、数据挖掘发现", 1)
keypoint(doc, "买方可分层经营(数百家高频金主作为重点客户);不同品类需匹配不同采购方式;文本与共现揭示可捆绑供货的品类。")
para(doc, "在统计之外,用聚类、交叉与文本挖掘进一步刻画买方结构与采购规律,为精准获客与品类选择提供依据。", after=8)
heading(doc, "5.1 买方机构四象限(谁值得重点经营)", 2)
add_image(doc, "11_org_scatter.png", width=5.6, caption="图6 买方机构四象限:发标频率 × 单标均值")
para(doc, "图中按「发标频率(横轴)× 单标均值(纵轴)」把买方分成四个象限。两个上方象限金额都高,但策略完全不同,"
          "这正是要点所在:", size=10, after=4)
takeaway(doc, "右上=核心金主(高频高额)。", "约 200 家州厅/大市政,既频繁发标又单标金额大;一次建立关系即可换来持续复购,投入产出最高,是建 KA 名单长期经营的首选。")
takeaway(doc, "左上=偶发大单(低频高额)。", "数量更多(近千家),单笔金额同样可观,但两年只发一两次标;它们值得「出现即投」抓单次大单,却不宜为其常设团队/渠道——因为缺乏复购、可预测性差,且部分高均值由一次性巨标拉高。")
takeaway(doc, "右下=跑量练兵、左下=长尾小标。", "右下高频小额适合积累中标信用;左下占比最大的小市政长尾性价比低,平均撒资源是浪费。")
para(doc, "换言之:右上看的是「关系的终身价值」,左上看的是「单次机会」——两者并非谁的金额更大,而是回报模式不同。下表为统计聚类对应的四类买方画像:", size=10, after=6)
cl = dm["org_clusters"]
rows = [[f"群{r['cluster']}", num(r["n_orgs"]), num(r["median_bids"]), brl(r["median_total"]), brl(r["median_avg"])] for r in cl]
table(doc, ["聚类", "机构数", "发标数中位", "累计额中位", "单标均值中位"], rows, widths=[0.8, 1.0, 1.2, 1.5, 1.5], font=9, align_right={1, 2, 3, 4})
para(doc, "")
heading(doc, "5.2 采购方式 × 维度 交叉", 2)
para(doc, "读图方式:行=采购方式,列=业务维度,格内数字与颜色深浅=该组合的招标数。用它回答「想做某个维度,主要要走哪种采购方式」。", size=10, after=4)
add_image(doc, "12_modality_x_dim.png", width=6.0, caption="图7 采购方式 × 业务维度交叉(招标数)")
def hcell(mod_kw, dim):
    mx = dm["modality_x_dimension"]
    for i, m in enumerate(mx["index"]):
        if all(k in m for k in mod_kw):
            return mx["data"][i][mx["columns"].index(dim)] if dim in mx["columns"] else None
    return None
pe_big = hcell(["Pregão", "Eletrôn"], "大宗商贸"); pe_med = hcell(["Pregão", "Eletrôn"], "医疗医药")
co_unc = hcell(["Concorrência", "Eletrôn"], "未分类"); cr_med = hcell(["Credenciamento"], "医疗医药")
para(doc, f"电子拍卖(Pregão Eletrônico)是绝对主通道,几乎在每个维度都最多——仅大宗商贸一格就有 {num(pe_big)} 条、医疗医药 {num(pe_med)} 条。"
          "含义:无论中企最终做哪个维度,都应先把电子拍卖的线上投标流程(电子签证、报价、缴保函)跑通,这是覆盖面最广的入口。", size=10, after=6)
para(doc, f"竞争性招标(Concorrência)集中在大额、复杂的工程与综合采购(如未分类 {num(co_unc)} 条),流程更正式、金额更大,适合具备总包或联合体能力后再切入;"
          f"认证入库(Credenciamento)在医疗服务相对突出(医疗医药 {num(cr_med)} 条),是该领域长期挂网、随到随审的通道。", size=10, after=6)
takeaway(doc, "先看维度走哪种方式,再备对应资质。", "进入某品类前,据其主流采购方式准备投标形式与资质,避免「做医疗却只会投电子拍卖、碰到认证入库无从下手」式的错配。")
heading(doc, "5.3 维度共现与品类语义", 2)
rows = [[r["维度对"], num(r["n"])] for r in dm["dimension_cooccurrence"][:8]]
table(doc, ["维度组合", "共现次数"], rows, widths=[3.5, 1.5], align_right={1})
kw = dm.get("keywords_by_dimension", {})
def kw_line(d):
    terms = [w["词"] for w in kw.get(d, [])][:8]
    return "、".join(terms) if terms else "—"
para(doc, "标的文本高频词揭示各维度真实采购内容(节选):", size=10, after=3)
for d in ["医疗医药", "高端制造", "数字经济", "大宗商贸"]:
    if d in kw:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(d + ":"); _font(r, 9.5, bold=True)
        r2 = p.add_run(kw_line(d)); _font(r2, 9.5)
para(doc, "医疗医药与大宗商贸高频共现(医用耗材兼具商贸属性),提示可做跨品类捆绑供货。", size=9.5, color=GREY, before=4)
page_break(doc)

# ════ 六、行动建议(附数据验证)════
heading(doc, "六、对中企的行动建议(附数据验证)", 1)
vlist = VAL["validations"]; recs = SY["recommendations"]
cnt = Counter(v["supported"] for v in vlist)
para(doc, f"以下建议均用本数据集逐条验证、且引用数字经独立审计核对:{cnt.get('数据支持',0)} 条「数据支持」、"
          f"{cnt.get('部分支持',0)} 条「部分支持」。「部分支持」指方向被数据印证,但执行假设(响应时效、转化率)或战略动作(并购本地公司)须靠外部尽调。", after=8)
SUPC = {"数据支持": GREEN, "部分支持": AMBER, "数据不足/需外部验证": RED}
n = min(len(recs), len(vlist))
for i in range(n):
    rec = recs[i]; v = vlist[i]
    p = doc.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(1)
    pr = rec.get("priority", "")
    if pr:
        rp = p.add_run(f"【{pr}】"); _font(rp, 10.5, bold=True, color=NAVY)
    ra = p.add_run(cx(rec["action"])); _font(ra, 10.5, bold=True)
    rv = p.add_run(f"  〔{v['supported']}〕"); _font(rv, 9.5, bold=True, color=RGBColor.from_string(SUPC.get(v["supported"], "595959")))
    pe = para(doc, "验证依据:" + v["evidence"], size=9.5, color=GREY, after=7, indent=0.3)
para(doc, "总体:" + VAL["overall"], size=9.5, color=GREY, before=2)
page_break(doc)

# ════ 附录:数据说明 ════
heading(doc, "附录:数据清洗与局限", 1)
rules = CLEAN["rules"]; b, a = CLEAN["before"], CLEAN["after"]
para(doc, "为保证分析可信,原始数据先行清洗:", after=4)
for t in [
    "剔除 2 笔确认录入错误(75.2亿 BRL 小镇公园招标、2.8万亿 BRL 合同,均为笔误);",
    f"年度采购计划删除 {rules['年度采购计划']['exact_full_row_duplicates_removed']:,} 条整行重复({b['年度采购计划']:,}→{a['年度采购计划']:,});",
    "金额=0 视为「预算未披露」(PNCP 允许保密)而非真零,仅在金额统计中排除、数量统计中保留;",
    f"已签合同标记 {rules['已签合同'].get('n_duration_suspect',0):,} 笔工期存疑、采购明细标记 {rules['采购明细'].get('n_item_no_unreliable',0):,} 行项号不可靠。",
]:
    p = doc.add_paragraph(style="List Bullet"); r = p.add_run(cx(t)); _font(r, 10); p.paragraph_format.space_after = Pt(3)
para(doc, "使用局限:", after=4, before=4)
for c in SY["caveats"][:5]:
    p = doc.add_paragraph(style="List Bullet"); r = p.add_run(cx(c[:150])); _font(r, 10); p.paragraph_format.space_after = Pt(3)

s = doc.settings.element; el = OxmlElement("w:updateFields"); el.set(qn("w:val"), "true"); s.append(el)
out = OUT / "巴西政府采购市场分析报告.docx"
doc.save(out)
print("✓", out, "| 段落", len(doc.paragraphs), "| 表格", len(doc.tables))
