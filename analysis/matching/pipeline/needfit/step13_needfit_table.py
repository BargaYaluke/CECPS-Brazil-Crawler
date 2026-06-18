# -*- coding: utf-8 -*-
"""Step13(需求匹配版·按匹配强度): 生成 中葡企业_巴西竞标匹配_三分类主表_需求匹配版.xlsx。
   口径: 纯需求契合(企业产品/业务+出海需求 ↔ 标的需要),不筛时效(过期标保留并标注)、
        不评估本地化与产能可行性。三类=匹配强度: ①直接对口(fit8-10) ②基本对口(fit6-7) ③同行业相关(fit4-5)。
   与 v3 同列同格式。链接: PNCP官方门户链接确定性构造(100%覆盖)+源站链接(DB)。
   ME/EPP 仅标注(本表不按其剔行,与"忽略本地化可行性"口径一致)。"""
import sys, io, json, glob, os, math, re
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parents[1]))
from match_lib import OUT, ROOT

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

comp = {c["cid"]: c for c in json.loads((OUT/"company_master.json").read_text("utf-8"))}
prof = json.loads((OUT/"company_profiles.json").read_text("utf-8"))
cat1 = json.loads((OUT/"cat1_needfit.json").read_text("utf-8"))
cat2 = json.loads((OUT/"cat2_needfit.json").read_text("utf-8"))
cat3 = json.loads((OUT/"cat3_needfit.json").read_text("utf-8"))

# 合并所有已有 PNCP 缓存(老+ v3),其余靠确定性构造/ DB 源站链接
ref, me = {}, {}
for f in ["pncp_links_refetch.json", "pncp_links_refetch_v3.json"]:
    p = OUT/f
    if p.exists(): ref.update(json.loads(p.read_text("utf-8")))
for f in ["meepp_benefit.json", "meepp_benefit_v3.json"]:
    p = OUT/f
    if p.exists(): me.update(json.loads(p.read_text("utf-8")))

L1 = "① 直接对口(强匹配)"
L2 = "② 基本对口(中匹配)"
L3 = "③ 同行业相关(弱匹配)"

t = pd.read_excel(ROOT/"data/exports/report.xlsx", sheet_name="招标主表")
deadline = {str(r["PNCP编号"]): r["投标截止"] for _, r in t.iterrows()}

PNCP_RE = re.compile(r"^(\d{14})-(\d+)-(\d+)/(\d{4})$")
def portal_of(pncp):
    m = PNCP_RE.match(str(pncp).strip())
    if not m: return ref.get(str(pncp), {}).get("portal")
    cnpj, _x, seq, ano = m.groups()
    return f"https://pncp.gov.br/app/editais/{cnpj}/{ano}/{int(seq)}"

def fmt_deadline(dt):
    if pd.isna(dt): return None
    dt = pd.to_datetime(dt)
    if (dt.hour, dt.minute) in [(0, 0), (23, 59)]:
        return dt.strftime("%Y-%m-%d")
    return dt.strftime("%Y-%m-%d %H:%M")

def deadline_cell(pncp, dias, prazo):
    suf = " ⚠已过期" if prazo == "已过期" else ""
    if dias is None or (isinstance(dias, float) and math.isnan(dias)):
        return "无截止"
    f = fmt_deadline(deadline.get(str(pncp)))
    return (f + suf) if f else (f"{int(dias)}天(采集日2026-05-29起)" + suf)

def fix_scheme(u):
    if u is None: return None
    u = str(u).strip()
    if not u or u.lower() == "nan": return None
    if u.lower().startswith(("http://", "https://")): return u
    if "." in u.split("/")[0]: return "https://" + u
    return None

def src_link(pncp, db_link):
    # 优先行级 DB 源站链接;回退缓存
    return fix_scheme(db_link) or fix_scheme(ref.get(str(pncp), {}).get("orig_link")) \
        or fix_scheme(ref.get(str(pncp), {}).get("linkSistemaOrigem"))

def link_status(pncp, has_src):
    r = ref.get(str(pncp))
    if r is not None and not r.get("api_ok"):
        return "⚠ 已从PNCP下架/不可查(以PNCP官方链接为准)"
    if has_src:
        return "源站链接(未逐一连通性检测);打不开请用PNCP官方链接"
    return "PNCP未登记源站链接 → 用PNCP官方链接"

comp_by_name = {c["公司简称"]: c for c in comp.values() if c.get("公司简称")}
def meepp_label(pncp, cn):
    info = me.get(str(pncp))
    if not info: return "", "none"          # 本表未逐一核查
    lab = info.get("label") or ""
    if "全部标项" in lab:
        return "⚠仅限ME/EPP(全部标项)·外资/大型通常不符资格,改JV或核实", "yel"
    if "部分标项" in lab:
        return lab.replace("ME/EPP专属", "仅限ME/EPP") + "·大企业仅可投其余标项", "yel"
    if "预留" in lab:
        return "含ME/EPP预留份额·大企业可投主份额", "none"
    if "分包" in lab:
        return "要求向ME/EPP分包", "yel"
    if "无ME/EPP" in lab:
        return "无ME/EPP限制", "none"
    return "", "none"

def sval(x):
    if x is None or isinstance(x, float): return None
    s = str(x).strip()
    return s or None

# 校验裁决合并
v1, v2 = {}, {}
for f in glob.glob(str(OUT/"cat_verdicts_nf"/"*.json")):
    try: d = json.loads(open(f, encoding="utf-8").read())
    except Exception: continue
    cid = d.get("cid") or os.path.splitext(os.path.basename(f))[0]
    for x in d.get("cat1", []):
        if x.get("pncp"): v1[f"{cid}|{x['pncp']}"] = x
    for x in d.get("cat2", []):
        if x.get("pncp"): v2[f"{cid}|{x['pncp']}"] = x
VERIFIED = bool(v1) or bool(v2)   # 无任何裁决文件 = 本数据集未做对抗校验

def imp(cid): return comp[cid]["重要程度"]
def imprank(cid): return comp[cid]["imp_rank"]

def cap_display(r):
    cap = sval(r.get("core_capability")) or ""
    gap = sval(r.get("capability_gap"))
    if gap and gap not in ("无", "未知"):
        return f"{cap}(⚠缺:{gap})" if cap else f"⚠缺:{gap}"
    return cap or None

def base_row(r, label, tier_n):
    cid = r["cid"]
    return {
        "类别": label, "_sort": (tier_n, imprank(cid)),
        "重要程度": imp(cid), "公司简称": comp[cid]["公司简称"],
        "行业": comp[cid].get("行业") or "(未分类)", "本地化状态": comp[cid]["本地化状态"],
        "出海需求": sval(comp[cid].get("出海需求")),
        "核心能力要求": cap_display(r),
        "标的类型": r.get("类型"), "竞标中文梗概": r.get("梗概"), "金额万CNY": r.get("金额万"),
        "采购方式": r.get("采购方式"), "提交截止日期": deadline_cell(r.get("pncp"), r.get("dias"), r.get("prazo")),
        "州": r.get("州"), "价格登记SRP": "是" if r.get("srp") else "否",
        "长期挂网": "是" if r.get("lt") else "否",
        "采购机构": r.get("机构"), "PNCP编号": r.get("pncp"), "_link": r.get("link"),
    }

rows = []
kept1 = drop1 = 0
for r in cat1:
    vd = v1.get(f"{r['cid']}|{r['pncp']}")
    if vd and vd.get("结论") == "剔除":
        drop1 += 1; continue
    kept1 += 1
    row = base_row(r, L1, 1)
    row["落地建议"] = r.get("action") or "对接采购方了解规格与供货要求"
    row["匹配理由"] = ((vd or {}).get("说明") or r.get("reason") or "") + ("" if VERIFIED else "(未对抗校验)")
    rows.append(row)
kept2 = drop2 = 0
for r in cat2:
    vd = v2.get(f"{r['cid']}|{r['pncp']}")
    if vd and vd.get("结论") == "不可行":
        drop2 += 1; continue
    kept2 += 1
    row = base_row(r, L2, 2)
    zizhi = sval((vd or {}).get("所需资质"))
    act = r.get("action") or "对接采购方了解需求"
    row["落地建议"] = (f"{act};涉及本地资质:【{zizhi}】" if zizhi else act)
    row["匹配理由"] = ((vd or {}).get("说明") or r.get("reason") or "") + ("" if VERIFIED else "(未对抗校验)")
    rows.append(row)
# ③ 弱匹配(未经对抗校验,作需求情报)
for r in cat3:
    row = base_row(r, L3, 3)
    row["落地建议"] = r.get("action") or "同行业线索,待核实具体规格是否对口"
    row["匹配理由"] = (r.get("reason") or "") + "(同行业弱匹配,未对抗校验)"
    rows.append(row)

# 链接 / ME-EPP / SRP
for row in rows:
    pncp = row["PNCP编号"]
    s = src_link(pncp, row.pop("_link", None))
    row["投标链接"] = s
    row["PNCP官方链接"] = portal_of(pncp)
    row["链接状态"] = link_status(pncp, bool(s))
    lab, flag = meepp_label(pncp, row["公司简称"])
    row["ME/EPP限制"] = lab
    row["_meepp_flag"] = flag
    row["SRP解释"] = None

COLS = ["类别","重要程度","公司简称","行业","本地化状态","出海需求","核心能力要求","标的类型","竞标中文梗概","金额万CNY",
        "采购方式","提交截止日期","州","价格登记SRP","长期挂网","落地建议","匹配理由","采购机构",
        "PNCP编号","投标链接","PNCP官方链接","链接状态","ME/EPP限制","SRP解释"]
SRP_TEXT = ('政府办一次招标,把中标供应商和报价"登记入库"(形成一份价格登记表 ata),'
            '之后在有效期内(通常 1 年)按需分批下单,不必每次重新招标;'
            '其他机构还能"搭车(carona)"用这份登记表采购。表里的 是/否 = 这个标走不走这个机制。')

df = pd.DataFrame(rows).sort_values(["_sort","公司简称","金额万CNY"], na_position="last")
flags = df["_meepp_flag"].tolist()
df = df.drop(columns=["_sort","_meepp_flag"])[COLS]
df.iloc[0, COLS.index("SRP解释")] = SRP_TEXT

WIDTHS = {"类别":18,"竞标中文梗概":50,"匹配理由":34,"落地建议":30,"采购机构":24,"采购方式":17,
          "投标链接":40,"PNCP编号":24,"公司简称":16,"行业":10,"本地化状态":10,"核心能力要求":24,
          "出海需求":46,"提交截止日期":18,"PNCP官方链接":42,"链接状态":30,"ME/EPP限制":28,"SRP解释":50}
hdr_fill = PatternFill("solid", fgColor="1F4E78"); hdr_font = Font(color="FFFFFF", bold=True, size=10)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cat_fill = {L1:"C6EFCE", L2:"FFF2CC", L3:"F2F2F2"}
imp_fill = {"1高":"FFC7CE","2中":"FFEB9C","3低":"D9E1F2"}
yel = PatternFill("solid", fgColor="FFF2CC")
VERIFY_NOTE = ("①②已逐条经独立智能体对抗校验(剔关键词假阳性/类型错配/需求错配)。"
               if VERIFIED else
               "⚠ 本表①②未经对抗校验(仅 DeepSeek 初评),可能含关键词假阳性/类型错配,投标前请人工复核。")
HDR_COMMENTS = {
    "类别": "本表为【需求匹配版·按匹配强度】:只按需求契合配对(企业产品/业务+出海需求 ↔ 标的需要),"
           "刻意不筛投标时效(已过期标保留并在截止列以『⚠已过期』标注)、不评估本地化与产能可行性。"
           "①直接对口=企业产品/服务就是标的所采购(DeepSeek fit 8-10);②基本对口=基本对口需小幅适配(fit 6-7);"
           "③同行业相关=同行业大类但产品不直接对口(fit 4-5,未经对抗校验,作需求情报)。"
           + VERIFY_NOTE,
    "出海需求": "来源:出海需求表(列『出海诉求/跟进记录』),按公司名连接,多条已去重合并。"
               "需求方向与『向巴西(政府)供货』错配的企业(从巴西进口/投资/考察等)已在评分与校验整体剔除。",
    "核心能力要求": "据梗概(非标题关键词)推断该标真正需要供方具备的核心履约能力:纯货物供应/工程总包/设备运维/持证人力/本地常驻劳务/持牌网络运营/本地配送等。"
                  "『⚠缺:X』=企业相对该能力的缺口。能力错配(行业/关键词相关但实际不能履约——如『医院』实为楼宇改造需工程总包、『打印机』实为品牌设备维保、『支付卡』实为封闭支付网络运营)的项已在评分阶段压低 fit、移出本表。",
    "提交截止日期": "取自源表(report.xlsx 招标主表)真实『投标截止』(巴西利亚时间)。"
                  "本表不筛时效:已过期标保留并以『⚠已过期』标注,作为需求情报;『无截止』为长期挂网/价格登记类。",
    "PNCP官方链接": "PNCP 官方门户页面(pncp.gov.br/app/editais/...),由 PNCP编号确定性构造,每个标的稳定权威入口,"
                  "100% 覆盖;源站链接打不开时,从此页可查看公告全文并跳转原招标系统。",
    "链接状态": "源站投标链接(取自采集库)未在本表逐一做连通性检测;多数巴西采购平台为 SPA/需登录/反爬,"
              "真人浏览器多可打开;打不开或无源站链接时,一律改用左侧『PNCP官方链接』。",
    "ME/EPP限制": "巴西 ME/EPP(中小企业)专属规则,仅标注、本表不据此剔行(与'忽略本地化/资格可行性'口径一致)。"
                "空白=本表未逐一核查该标 ME/EPP 状态;如需投标须自行核验标项级 tipoBenefício。"
                "据 LC123/2006 Art.3 §4,外资分支/有外资法人股东的子公司通常不具 ME/EPP 资格。",
    "SRP解释": "SRP=Sistema de Registro de Preços(价格登记制),见本列首行说明。",
}

LABEL = os.environ.get("MATCH_LABEL", "")   # 数据集后缀,如 _在巴中企;默认空=原 437 家
out_path = ROOT/f"data/exports/中葡企业_巴西竞标匹配_三分类主表_需求匹配版{LABEL}.xlsx"
with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
    df.to_excel(xw, sheet_name="竞标匹配主表(三分类)", index=False)
    ws = xw.book["竞标匹配主表(三分类)"]
    hdrs = {ws.cell(1,ci).value: ci for ci in range(1, ws.max_column+1)}
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(df)+1}"
    for cell in ws[1]:
        cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = hdr_align
    for name, txt in HDR_COMMENTS.items():
        ws.cell(1, hdrs[name]).comment = Comment(txt, "Claude")
    for ci in range(1, len(COLS)+1):
        ws.column_dimensions[get_column_letter(ci)].width = WIDTHS.get(ws.cell(1,ci).value, 12)
    for ri in range(2, len(df)+2):
        for ci in range(1, len(COLS)+1):
            ws.cell(ri,ci).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(ri,ci).font = Font(size=9)
        cv = str(ws.cell(ri, hdrs["类别"]).value or "")
        if cv in cat_fill:
            ws.cell(ri, hdrs["类别"]).fill = PatternFill("solid", fgColor=cat_fill[cv])
        iv = str(ws.cell(ri, hdrs["重要程度"]).value or "")
        if iv in imp_fill:
            ws.cell(ri, hdrs["重要程度"]).fill = PatternFill("solid", fgColor=imp_fill[iv])
        if flags[ri-2] == "yel":
            ws.cell(ri, hdrs["ME/EPP限制"]).fill = yel

from collections import Counter
print("已生成:", out_path)
print("行数:", len(df))
print("分类分布:", dict(Counter(df["类别"])))
print(f"① 校验保留{kept1}/剔除{drop1} | ② 保留{kept2}/剔除{drop2}")
print("涉及企业数:", df["公司简称"].nunique())
print("含已过期标行数:", int(df["提交截止日期"].astype(str).str.contains("已过期").sum()))
print("PNCP官方链接覆盖:", int(df["PNCP官方链接"].notna().sum()), "/", len(df))
print("投标链接(源站)非空:", int(df["投标链接"].notna().sum()), "/", len(df))
