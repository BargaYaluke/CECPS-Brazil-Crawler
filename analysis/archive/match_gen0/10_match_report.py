# -*- coding: utf-8 -*-
"""
10_match_report.py — 生成《中葡企业 × 巴西政府采购 竞标匹配策略报告》Word
依赖 09 产出的 analysis/outputs/match/consolidated.json
"""
import os, json, collections, datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = r"C:\巴西爬虫"
OUT  = os.path.join(ROOT, "analysis", "outputs", "match")
DOCX_OUT = os.path.join(ROOT, "analysis", "outputs", "中葡企业_巴西竞标匹配策略报告.docx")
TODAY = datetime.date(2026, 6, 9)
DIM2ZH = {'high_end_manufacturing':'高端制造','healthcare':'医疗医药','digital_economy':'数字经济',
 'commodities':'大宗商贸','culture_sports':'文化体育','cross_border_ecommerce':'跨境商贸','None':'未分类'}
NAVY = RGBColor(0x1F,0x4E,0x78)

data = json.load(open(os.path.join(OUT,'consolidated.json'),encoding='utf-8'))
rows, overview, missing = data['rows'], data['overview'], data['missing']
by_idx_recs = collections.defaultdict(list)
for r in rows: by_idx_recs[r['short']].append(r)

doc = Document()
st = doc.styles['Normal']; st.font.name='微软雅黑'; st.font.size=Pt(10.5)
try:
    st._element.rPr.rFonts.set(__import__('docx').oxml.ns.qn('w:eastAsia'),'微软雅黑')
except Exception: pass

def H(text, lvl=1):
    h=doc.add_heading(text, level=lvl)
    for run in h.runs: run.font.color.rgb=NAVY
    return h
def P(text, size=10.5, bold=False, color=None):
    p=doc.add_paragraph(); r=p.add_run(text); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=color
    return p
def money(v): return f"{(v or 0)/10000:.0f}万" if (v or 0)>=10000 else f"{(v or 0)/10000:.1f}万"

# ---------- 封面
t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=t.add_run('中葡企业 × 巴西政府采购\n竞标匹配策略报告'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=NAVY
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
sr=sub.add_run(f'基于 PNCP 在投招标 × 438家意向出海企业 的数据匹配\n生成日期 {TODAY}　时效基准日 {TODAY}')
sr.font.size=Pt(11); sr.font.color.rgb=RGBColor(0x60,0x60,0x60)
doc.add_paragraph()

# ---------- 执行摘要
H('一、执行摘要',1)
fit_cnt=collections.Counter(o['gov_fit'] for o in overview)
imp_cnt=collections.Counter(o['importance'] for o in overview)
n_strong=sum(1 for o in overview if o['gov_fit']=='强')
P(f"本报告对 438 家有意出海巴西的中葡企业，逐一从 PNCP 在投政府采购竞标中做了语义匹配、可行性筛选与时效分层，"
  f"共产出企业-竞标推荐 {len(rows)} 条，覆盖 {len(overview)} 家企业。", bold=True)
P("核心结论：")
bullets=[
 f"政采契合度分布：强 {fit_cnt.get('强',0)} 家、中 {fit_cnt.get('中',0)} 家、弱 {fit_cnt.get('弱',0)} 家、几乎无 {fit_cnt.get('几乎无',0)} 家。"
 "契合度高的集中在医疗医药(药品/耗材/设备)、高端制造(车辆/工程机械/电力)、数字经济(安防/IT设备)三条主线。",
 "本地化是中标的硬门槛：纯境外身份直接中标≈0(全样本仅15笔/0.04%)。所有推荐均以"
 "『先取得巴西法人(PJ)身份或与本地PJ联合体/收购』为前提，绝非中标后再补。",
 "时效是第一道闸：在投标的约九成将在14天内截止，未本地化企业无法直投，须走"
 "『与本地企业组联合体』或转向年度计划(PCA)提前3-6个月卡位、长期挂网(SRP/Credenciamento)从容布局。",
 "可行性优先：只匹配中企能『供货』的实物品类，已剔除需本地施工资质的工程(obra)与本地人力服务"
 "(维修/运输/清洁/安保/培训/金融/医疗人力)；并按企业规模匹配金额量级，远超产能的大单建议联合体或降级。",
 "诚实提示：教育培训、纯电商零售、消费电子(如E-bike)等以服务/零售为主的企业，政府采购并非其主战场，"
 "本报告如实标注其契合度为弱，并给出更现实的替代路径(B2B分销/PCA/长期挂网)。",
]
for b in bullets:
    p=doc.add_paragraph(style='List Bullet'); p.add_run(b).font.size=Pt(10.5)

# ---------- 方法论
H('二、匹配方法论与口径',1)
steps=[
 ('数据来源','招标侧=巴西PNCP门户(采集窗口2026-05-22~05-28，8357笔在投/可投标的)；企业侧=中葡企业机构数据库438家(公司全称/六大行业/重要程度/主营业务)。'),
 ('数据清洗','统一行业→PNCP六维度映射；招标按时效基准日重算剩余天数；翻译缓存回填中文梗概(覆盖>99%标的)。'),
 ('可行性过滤','以"供货动作词"(aquisição/fornecimento)判定供货标，强制剔除工程(obra/construção/reforma)与本地人力服务(维修/运输/清洁/安保/培训/金融/医疗人力/带操作员租赁/燃油)等中企无法承接者。'),
 ('双语子类目桥接','中文主营业务与葡文标的之间，用药品/医疗设备/耗材/车辆/新能源充电/工程机械/电力设备/IT软硬件/安防监控/体育器材等子类目做跨语言对接，生成候选池。'),
 ('大模型语义匹配','按重要程度分层，对每家企业用大模型从候选池做跨语义匹配+产能判断+时效渠道+理由，剔除子类目误配的噪声；对全部1高企业再做独立对抗校验，剔除不现实匹配。'),
 ('排序原则','优先级=企业重要程度(1高>2中>3低>无标识) × 匹配质量；越重要的企业优先匹配越优质(契合度高、可行、金额匹配)的竞标。'),
]
tb=doc.add_table(rows=1,cols=2); tb.style='Light Grid Accent 1'; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
tb.rows[0].cells[0].text='环节'; tb.rows[0].cells[1].text='做法'
for k,v in steps:
    c=tb.add_row().cells; c[0].text=k; c[1].text=v
for cell in tb.rows[0].cells:
    for p in cell.paragraphs:
        for r in p.runs: r.bold=True

# ---------- 核心约束
H('三、四条匹配铁律(决定可行性)',1)
laws=[
 ('本地化分水岭','取得巴西法人身份(PJ)或与本地PJ深度合资/收购，是参与与中标的前提；可优先收购一家有CNPJ与政府采购履约记录的小型本地供应商以继承资质。'),
 ('双时间轴','『现在能投』(近期标，多在14天内截止) 与 『未来管道』(PCA年度计划) 同时打开；本地化需3-6个月，未本地化者近期标只能联合体或转PCA/长期挂网。'),
 ('只做供货、不碰工程与本地服务','聚焦实物供货(设备/药品/车辆/耗材)，规避需本地施工资质的Obras与本地人力服务。'),
 ('量力而行','按企业规模匹配金额量级；小企业不接远超产能的大单，可走联合体/分包；金额远超产能者降级或剔除。'),
]
for i,(k,v) in enumerate(laws,1):
    p=doc.add_paragraph(); p.add_run(f'{i}. {k}：').bold=True; p.add_run(v)

# ---------- 分层结果: 1高
H('四、重点企业匹配(1高，按企业)',1)
P("以下为重要程度最高(1高)企业的精选推荐(经对抗校验)。完整明细见配套Excel《竞标匹配总表》。", color=RGBColor(0x60,0x60,0x60))
highs=[o for o in overview if o['rank']==4]
highs.sort(key=lambda o:(o['gov_fit']!='强', o['dim'], o['short']))
for o in highs:
    recs=by_idx_recs.get(o['short'],[])[:4]
    if not recs and o['gov_fit'] in ('弱','几乎无'):
        H(f"{o['short']}（{DIM2ZH.get(o['dim'],o['dim'])}｜契合度{o['gov_fit']}）",3)
        P("政采契合度弱："+o['fit_reason'], size=10)
        if o['future_pipeline']: P("替代路径："+o['future_pipeline'], size=10)
        continue
    H(f"{o['short']}（{DIM2ZH.get(o['dim'],o['dim'])}｜契合度{o['gov_fit']}）",3)
    if o['capacity_assessment']: P("产能与定位："+o['capacity_assessment'], size=10)
    if recs:
        tb=doc.add_table(rows=1,cols=5); tb.style='Light List Accent 1'
        hdr=['竞标(中文梗概)','金额','截止/剩余','时效渠道','匹配理由']
        for j,h in enumerate(hdr):
            tb.rows[0].cells[j].text=h
            for p in tb.rows[0].cells[j].paragraphs:
                for r in p.runs: r.bold=True; r.font.size=Pt(9)
        for rc in recs:
            c=tb.add_row().cells
            c[0].text=rc['objeto_zh'][:60]
            c[1].text=money(rc['valor_cny'])
            c[2].text=f"{rc['deadline'] or '-'}\n剩{rc['dias']}天" if rc['dias'] is not None else (rc['deadline'] or '-')
            c[3].text=rc['channel'][:18]
            c[4].text=rc['rationale'][:90]
            for j in range(5):
                for p in c[j].paragraphs:
                    for r in p.runs: r.font.size=Pt(8.5)
    if o['future_pipeline']: P("未来管道(PCA)："+o['future_pipeline'][:200], size=9.5, color=RGBColor(0x60,0x60,0x60))

# ---------- 2中 概览
H('五、次重点企业匹配(2中，概览)',1)
mids=[o for o in overview if o['rank']==3]
P(f"共 {len(mids)} 家。下表列出各企业政采契合度与首选推荐方向，完整推荐见Excel。", color=RGBColor(0x60,0x60,0x60))
tb=doc.add_table(rows=1,cols=4); tb.style='Light Grid Accent 1'
for j,h in enumerate(['公司','行业','契合度','首选推荐(梗概/金额/渠道)']):
    tb.rows[0].cells[j].text=h
    for p in tb.rows[0].cells[j].paragraphs:
        for r in p.runs: r.bold=True; r.font.size=Pt(9)
for o in sorted(mids,key=lambda x:(x['gov_fit']!='强',x['dim'])):
    recs=by_idx_recs.get(o['short'],[])
    top=recs[0] if recs else None
    c=tb.add_row().cells
    c[0].text=o['short']; c[1].text=DIM2ZH.get(o['dim'],o['dim']); c[2].text=o['gov_fit']
    c[3].text=(f"{top['objeto_zh'][:34]}｜{money(top['valor_cny'])}｜{top['channel'][:14]}" if top else "(候选中无强匹配，建议转PCA/长期挂网)")
    for j in range(4):
        for p in c[j].paragraphs:
            for r in p.runs: r.font.size=Pt(8.5)

# ---------- 按行业机会图谱
H('六、按行业的机会图谱(含3低/无标识)',1)
P("3低与无标识企业(共323家)以行业模板批量匹配；下表汇总各行业可切入的典型品类与建议渠道。", color=RGBColor(0x60,0x60,0x60))
dim_rows=collections.defaultdict(list)
for r in rows: dim_rows[r['dim_zh']].append(r)
tb=doc.add_table(rows=1,cols=4); tb.style='Light Grid Accent 1'
for j,h in enumerate(['行业','企业数','可切入典型品类(来自实配标的)','主推渠道']):
    tb.rows[0].cells[j].text=h
    for p in tb.rows[0].cells[j].paragraphs:
        for r in p.runs: r.bold=True; r.font.size=Pt(9)
dim_comp=collections.Counter(o['dim'] for o in overview)
CHANNEL={'医疗医药':'电子竞价Pregão+SRP价格登记；ANVISA注册先行','高端制造':'供货+本地安装服务商；大单组联合体','数字经济':'Pregão电子标+长期挂网；设备+集成','大宗商贸':'SRP价格登记滚动供货','文化体育':'器材供货为主，避开活动/场馆工程','跨境商贸':'政采契合弱，建议B2B分销/制服等实物供货','未分类':'按主营再归类'}
for dimzh in ['医疗医药','高端制造','数字经济','大宗商贸','文化体育','跨境商贸']:
    rs=dim_rows.get(dimzh,[])
    cats=collections.Counter()
    for r in rs:
        cats[r['product_match'][:10] or '其它']+=1
    ncomp=sum(v for k,v in dim_comp.items() if DIM2ZH.get(k)==dimzh)
    sample=[r['objeto_zh'][:14] for r in rs[:3]]
    c=tb.add_row().cells
    c[0].text=dimzh; c[1].text=str(ncomp)
    c[2].text='；'.join(sample) if sample else '(候选有限)'
    c[3].text=CHANNEL.get(dimzh,'')
    for j in range(4):
        for p in c[j].paragraphs:
            for r in p.runs: r.font.size=Pt(8.5)

# ---------- 弱契合诚实清单
H('七、政采弱契合企业(诚实提示)',1)
weak=[o for o in overview if o['gov_fit'] in ('弱','几乎无')]
P(f"共 {len(weak)} 家企业的产品/服务与巴西政府采购契合度弱(多为教育培训、纯电商零售、消费品、咨询服务等)。"
  "对这些企业，政府采购不应作为出海主渠道，建议转向B2B分销、商超零售或PCA长周期布局。代表性企业：", )
for o in sorted(weak,key=lambda x:(-x['rank']))[:25]:
    p=doc.add_paragraph(style='List Bullet')
    p.add_run(f"{o['short']}（{DIM2ZH.get(o['dim'],o['dim'])}，{o['importance']}）：").bold=True
    p.add_run((o['fit_reason'] or '')[:80]).font.size=Pt(9.5)

# ---------- 局限
H('八、局限与使用建议',1)
for b in [
 "招标侧为2026-05-22~05-28单周快照，不可年化；金额为预估值，最终以标书为准。",
 "本地化周期(3-6月)、投标响应时效、中标率等执行假设需结合外部尽调；子类目为启发式，个别误配以大模型复核与人工二次确认为准。",
 "建议用法：1高/2中企业按本报告+Excel逐家落实本地化与联合体；3低/无标识按行业图谱批量触达；所有近期标先核验截止日与本地资质再投。",
]:
    p=doc.add_paragraph(style='List Bullet'); p.add_run(b).font.size=Pt(10)

doc.save(DOCX_OUT)
print("Word ->", DOCX_OUT, "| 推荐行", len(rows), "| 企业", len(overview), "| 弱契合", len(weak))
