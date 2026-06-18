# -*- coding: utf-8 -*-
"""Step9(纯契合版): 生成 中葡企业_巴西竞标匹配_三分类主表_纯契合版.xlsx。
   与既有主表同格式;不考虑时效/本地化可行性,新增『时效状态』列标注快照状态。
   三分类: ①产品直接对口(货物) ②业务相邻(需本地资质/服务类) ③高价值候选(不限企业)。"""
import sys, io, json, glob, os, math
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from match_lib import OUT, ROOT

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

comp = {c["cid"]: c for c in json.loads((OUT/"company_master.json").read_text("utf-8"))}
prof = json.loads((OUT/"company_profiles.json").read_text("utf-8"))
cat1 = json.loads((OUT/"cat1_pf.json").read_text("utf-8"))
cat2 = json.loads((OUT/"cat2_pf.json").read_text("utf-8"))
cat3 = json.loads((OUT/"cat3_pf.json").read_text("utf-8"))
ref  = json.loads((OUT/"pncp_links_refetch_pf.json").read_text("utf-8"))
conn = json.loads((OUT/"link_connectivity_pf.json").read_text("utf-8"))
me   = json.loads((OUT/"meepp_benefit_pf.json").read_text("utf-8"))

L1 = "① 产品直接对口(货物)"
L2 = "② 业务相邻(需本地资质/服务类)"
L3 = "③ 高价值候选(不限企业)"

t = pd.read_excel(ROOT/"data/exports/report.xlsx", sheet_name="招标主表")
deadline = {str(r["PNCP编号"]): r["投标截止"] for _, r in t.iterrows()}

def fmt_deadline(dt):
    if pd.isna(dt):
        return None
    dt = pd.to_datetime(dt)
    if (dt.hour, dt.minute) in [(0, 0), (23, 59)]:
        return dt.strftime("%Y-%m-%d")
    return dt.strftime("%Y-%m-%d %H:%M")

def deadline_cell(pncp, dias):
    if dias is None or (isinstance(dias, float) and math.isnan(dias)):
        return "无截止"
    f = fmt_deadline(deadline.get(str(pncp)))
    return f if f else f"{int(dias)}天(采集日2026-05-29起)"

def prazo_cell(r):
    dias = r.get("dias")
    if dias is None or (isinstance(dias, float) and math.isnan(dias)):
        return "无截止"
    s = r.get("时效状态")
    return s if s and not (isinstance(s, float) and math.isnan(s)) else "未知"

def fix_scheme(u):
    if u is None: return None
    u = str(u).strip()
    if not u or u.lower() == "nan": return None
    if u.lower().startswith(("http://", "https://")): return u
    if "." in u.split("/")[0]:
        return "https://" + u
    return None

def best_origin(pncp):
    r = ref.get(pncp, {})
    return fix_scheme(r.get("orig_link")) or fix_scheme(r.get("linkSistemaOrigem"))

def link_status(pncp):
    r = ref.get(pncp, {}); c = conn.get(pncp, {})
    if not r.get("api_ok"):
        return "⚠ 已从PNCP下架/不可查(可能已撤销或结束)"
    if not best_origin(pncp):
        return "PNCP未登记源站链接 → 用PNCP官方链接"
    cls = c.get("origin_class", "")
    if "HTTP可达" in cls or "200" in cls: return "源站链接可访问"
    if "重定向" in cls:                    return "源站可访问(有跳转)"
    if "403" in cls:                       return "源站反爬:脚本被拒,浏览器可正常打开(或用PNCP官方链接)"
    if "失效" in cls:                      return "源站链接已失效 → 用PNCP官方链接"
    return "源站暂时无法连接 → 用PNCP官方链接"

comp_by_name = {c["公司简称"]: c for c in comp.values() if c.get("公司简称")}
def is_large(cn):
    c = comp_by_name.get(cn, {})
    tier = (prof.get(c.get("cid"), {}) or {}).get("capacity_tier")
    return bool(tier) and tier >= 3

def meepp_label(pncp, cn):
    info = me.get(str(pncp), {})
    lab = info.get("label") or ""
    if "全部标项" in lab:
        if is_large(cn):
            return "仅限ME/EPP(全部标项)·大型企业不可投", "remove_large"
        return "⚠仅限ME/EPP(全部标项)·外资子公司通常不符资格,建议核实或改JV", "flag"
    if "部分标项" in lab:
        return lab.replace("ME/EPP专属", "仅限ME/EPP") + "·大企业仅可投其余标项", "flag"
    if "预留" in lab:
        return "含ME/EPP预留份额·大企业可投主份额", "ok"
    if "分包" in lab:
        return "要求向ME/EPP分包", "flag"
    return "", "none"

def sval(x):
    if x is None or isinstance(x, float):
        return None
    s = str(x).strip()
    return s or None

v1, v2 = {}, {}
for f in glob.glob(str(OUT/"cat_verdicts_pf"/"*.json")):
    try: d = json.loads(open(f, encoding="utf-8").read())
    except Exception: continue
    cid = d.get("cid") or os.path.splitext(os.path.basename(f))[0]
    for x in d.get("cat1", []):
        if x.get("pncp"): v1[f"{cid}|{x['pncp']}"] = x
    for x in d.get("cat2", []):
        if x.get("pncp"): v2[f"{cid}|{x['pncp']}"] = x

def imp(cid): return comp[cid]["重要程度"]
def imprank(cid): return comp[cid]["imp_rank"]

rows = []
kept1 = drop1 = 0
for r in cat1:
    vd = v1.get(f"{r['cid']}|{r['pncp']}")
    if vd and vd.get("结论") == "剔除":
        drop1 += 1; continue
    kept1 += 1
    rows.append({
        "类别": L1, "_sort": (1, imprank(r["cid"]), -(r.get("fit") or 0)),
        "重要程度": imp(r["cid"]), "公司简称": comp[r["cid"]]["公司简称"],
        "行业": comp[r["cid"]].get("行业") or "(未分类)", "本地化状态": comp[r["cid"]]["本地化状态"],
        "出海需求": sval(comp[r["cid"]].get("出海需求")),
        "标的类型": r.get("类型"), "竞标中文梗概": r.get("梗概"), "金额万CNY": r.get("金额万"),
        "采购方式": r.get("采购方式"), "提交截止日期": deadline_cell(r.get("pncp"), r.get("dias")),
        "时效状态": prazo_cell(r),
        "州": r.get("州"), "价格登记SRP": "是" if r.get("srp") else "否",
        "长期挂网": "是" if r.get("lt") else "否",
        "落地建议": r.get("channel"), "匹配理由": r.get("reason"),
        "采购机构": r.get("机构"), "PNCP编号": r.get("pncp"),
    })
kept2 = drop2 = 0
for r in cat2:
    vd = v2.get(f"{r['cid']}|{r['pncp']}")
    if vd and vd.get("结论") == "不可行":
        drop2 += 1; continue
    kept2 += 1
    zizhi = (vd or {}).get("所需资质") or "相应本地牌照/资质"
    note = (vd or {}).get("说明") or r.get("reason")
    rows.append({
        "类别": L2, "_sort": (2, imprank(r["cid"]), -(r.get("fit") or 0)),
        "重要程度": imp(r["cid"]), "公司简称": comp[r["cid"]]["公司简称"],
        "行业": comp[r["cid"]].get("行业") or "(未分类)", "本地化状态": comp[r["cid"]]["本地化状态"],
        "出海需求": sval(comp[r["cid"]].get("出海需求")),
        "标的类型": r.get("类型"), "竞标中文梗概": r.get("梗概"), "金额万CNY": r.get("金额万"),
        "采购方式": r.get("采购方式"), "提交截止日期": deadline_cell(r.get("pncp"), r.get("dias")),
        "时效状态": prazo_cell(r),
        "州": r.get("州"), "价格登记SRP": "是" if r.get("srp") else "否",
        "长期挂网": "是" if r.get("lt") else "否",
        "落地建议": f"需取得【{zizhi}】后可承接", "匹配理由": note,
        "采购机构": r.get("机构"), "PNCP编号": r.get("pncp"),
    })
for r in cat3:
    rows.append({
        "类别": L3, "_sort": (3, 0, -(r.get("金额万") or 0)),
        "重要程度": "", "公司简称": "—(高价值候选,不限企业)",
        "行业": r.get("行业"), "本地化状态": "", "出海需求": None,
        "标的类型": "", "竞标中文梗概": r.get("梗概"), "金额万CNY": r.get("金额万"),
        "采购方式": r.get("采购方式"), "提交截止日期": deadline_cell(r.get("pncp"), r.get("dias")),
        "时效状态": prazo_cell(r),
        "州": r.get("州"), "价格登记SRP": "是" if r.get("srp") else "否",
        "长期挂网": "是" if r.get("lt") else "否",
        "落地建议": "高价值大单:可组联合体(consórcio)/本地伙伴切入",
        "匹配理由": "金额高,该机构存在此类大额需求,值得跟踪", "采购机构": r.get("机构"),
        "PNCP编号": r.get("pncp"),
    })

for row in rows:
    pncp = row["PNCP编号"]
    row["投标链接"] = best_origin(pncp)
    row["PNCP官方链接"] = ref.get(pncp, {}).get("portal")
    row["链接状态"] = link_status(pncp)
    lab, action = meepp_label(pncp, row["公司简称"])
    row["ME/EPP限制"] = lab
    row["_meepp_action"] = action
    row["SRP解释"] = None

main_rows = [r for r in rows if r["_meepp_action"] != "remove_large"]
removed_rows = [r for r in rows if r["_meepp_action"] == "remove_large"]

COLS = ["类别","重要程度","公司简称","行业","本地化状态","出海需求","标的类型","竞标中文梗概","金额万CNY",
        "采购方式","提交截止日期","时效状态","州","价格登记SRP","长期挂网","落地建议","匹配理由","采购机构",
        "PNCP编号","投标链接","PNCP官方链接","链接状态","ME/EPP限制","SRP解释"]
SRP_TEXT = ('政府办一次招标,把中标供应商和报价"登记入库"(形成一份价格登记表 ata),'
            '之后在有效期内(通常 1 年)按需分批下单,不必每次重新招标;'
            '其他机构还能"搭车(carona)"用这份登记表采购。表里的 是/否 = 这个标走不走这个机制。')

df = pd.DataFrame(main_rows).sort_values(["_sort","公司简称"]).drop(columns=["_sort","_meepp_action"])[COLS]
dfr = (pd.DataFrame(removed_rows).sort_values(["_sort","公司简称"]).drop(columns=["_sort","_meepp_action"])[COLS]
       if removed_rows else pd.DataFrame(columns=COLS))
df.iloc[0, COLS.index("SRP解释")] = SRP_TEXT

WIDTHS = {"类别":22,"竞标中文梗概":50,"匹配理由":34,"落地建议":30,"采购机构":24,"采购方式":17,
          "投标链接":40,"PNCP编号":24,"公司简称":16,"行业":10,"本地化状态":10,"时效状态":10,
          "出海需求":46,"提交截止日期":16,"PNCP官方链接":40,"链接状态":30,"ME/EPP限制":30,"SRP解释":50}
hdr_fill = PatternFill("solid", fgColor="1F4E78"); hdr_font = Font(color="FFFFFF", bold=True, size=10)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cat_fill = {L1:"C6EFCE", L2:"FFF2CC", L3:"DDEBF7"}
imp_fill = {"1高":"FFC7CE","2中":"FFEB9C","3低":"D9E1F2"}
exp_fill = PatternFill("solid", fgColor="D9D9D9")
yel = PatternFill("solid", fgColor="FFF2CC"); red = PatternFill("solid", fgColor="FFC7CE")
HDR_COMMENTS = {
    "类别": "纯契合三分类(本表不考虑时效与本地化可行性,只看需求匹配):"
           "①标的货物与企业产品直接对口;②与企业业务相邻的服务/资质类机会;③高价值大单(不限企业)。",
    "出海需求": "来源:中葡企业机构数据库_企业拜访表(列『出海诉求/跟进记录』),按公司名连接,去重合并。"
               "匹配时已校验需求方向:与『向巴西政府供货』错配(纯进口/引进/投资并购/仅考察)的企业已剔除。",
    "提交截止日期": "取自源表(report.xlsx 招标主表)真实『投标截止』,巴西利亚当地时间;"
                  "本表不按时效筛选,已过期的标也保留(见『时效状态』列)。",
    "时效状态": "采集快照(2026-05-29)时的状态:还能投/临近截止/已过期/无截止。"
              "本表为纯契合版,已过期标仍保留 = 该机构存在此类采购需求,作需求情报与下一轮同类标跟踪用。",
    "PNCP官方链接": "PNCP 官方门户页面(pncp.gov.br),每个标的稳定权威入口;"
                  "源站(投标链接)打不开时,从此页可查看公告全文并跳转到原招标系统。",
    "链接状态": "源站投标链接的连通性(自动检测)。注:多数巴西采购平台为 SPA/需登录/反爬,"
              "'反爬'类链接在真人浏览器中可正常打开;'已下架'多为已结束/已撤销的标(本表保留作需求情报)。",
    "ME/EPP限制": "巴西 ME/EPP(中小企业)专属规则。'仅限ME/EPP(全部标项)'=整标只许巴西中小企业投,"
                "大型企业不可投;'部分标项'=仅部分标项保留给中小企业。"
                "据 LC123/2006 Art.3 §4,外资分支/有外资法人股东的子公司通常不具 ME/EPP 资格。",
    "SRP解释": "SRP=Sistema de Registro de Preços(价格登记制),见本列首行说明。",
}

def style_sheet(ws, n_rows, hdrs):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n_rows+1}"
    for cell in ws[1]:
        cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = hdr_align
    for name, txt in HDR_COMMENTS.items():
        ws.cell(1, hdrs[name]).comment = Comment(txt, "Claude")
    for ci in range(1, len(COLS)+1):
        ws.column_dimensions[get_column_letter(ci)].width = WIDTHS.get(ws.cell(1,ci).value, 12)
    for ri in range(2, n_rows+2):
        for ci in range(1, len(COLS)+1):
            ws.cell(ri,ci).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(ri,ci).font = Font(size=9)
        cv = str(ws.cell(ri, hdrs["类别"]).value or "")
        if cv in cat_fill:
            ws.cell(ri, hdrs["类别"]).fill = PatternFill("solid", fgColor=cat_fill[cv])
        iv = str(ws.cell(ri, hdrs["重要程度"]).value or "")
        if iv in imp_fill:
            ws.cell(ri, hdrs["重要程度"]).fill = PatternFill("solid", fgColor=imp_fill[iv])
        tv = str(ws.cell(ri, hdrs["时效状态"]).value or "")
        if tv == "已过期":
            ws.cell(ri, hdrs["时效状态"]).fill = exp_fill
        mv = str(ws.cell(ri, hdrs["ME/EPP限制"]).value or "")
        if mv.startswith("⚠") or "部分标项" in mv or "分包" in mv:
            ws.cell(ri, hdrs["ME/EPP限制"]).fill = yel
        elif "大型企业不可投" in mv:
            ws.cell(ri, hdrs["ME/EPP限制"]).fill = red

out_path = ROOT/"data/exports/中葡企业_巴西竞标匹配_三分类主表_纯契合版.xlsx"
with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
    df.to_excel(xw, sheet_name="竞标匹配主表(纯契合三分类)", index=False)
    ws = xw.book["竞标匹配主表(纯契合三分类)"]
    hdrs = {ws.cell(1,ci).value: ci for ci in range(1, ws.max_column+1)}
    style_sheet(ws, len(df), hdrs)
    if len(dfr):
        dfr.to_excel(xw, sheet_name="已剔除_ME-EPP专属(大型企业)", index=False)
        ws2 = xw.book["已剔除_ME-EPP专属(大型企业)"]
        hdrs2 = {ws2.cell(1,ci).value: ci for ci in range(1, ws2.max_column+1)}
        style_sheet(ws2, len(dfr), hdrs2)

from collections import Counter
print("已生成:", out_path)
print("主表行数:", len(df), "| ME/EPP剔除行:", len(dfr))
print("分类分布:", dict(Counter(df["类别"])))
print(f"第一类 校验保留{kept1}/剔除{drop1} | 第二类 保留{kept2}/剔除{drop2}")
print("涉及企业数:", df.loc[df["类别"]!=L3,"公司简称"].nunique())
print("时效状态分布:", dict(Counter(df["时效状态"])))
print("出海需求非空行:", int(df["出海需求"].notna().sum()))
