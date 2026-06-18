"""把企业拜访表的『出海需求/诉求』并入竞标匹配三分类主表（按公司名连接）。"""
import copy
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

VISIT = "data/exports/中葡企业机构数据库_企业拜访表_企业拜访表.xlsx"
TARGET = "data/exports/中葡企业_巴西竞标匹配_三分类主表.xlsx"
SRC_COL = "出海诉求/跟进记录"
NEW_HEADER = "出海需求"
INSERT_AT = 6  # 插在 F 列（本地化状态之后、标的类型之前）

def clean(s):
    s = str(s).strip()
    return "" if s.lower() in ("", "nan", "none") else s

# 1) 拜访表 -> {企业名称: 合并后的出海诉求}
v = pd.read_excel(VISIT, dtype=str)
v["企业名称"] = v["企业名称"].astype(str).str.strip()
need = {}
for name, grp in v.groupby("企业名称"):
    vals = []
    for x in grp[SRC_COL]:
        c = clean(x)
        if c and c not in vals:
            vals.append(c)
    if vals:
        need[name] = "\n\n".join(vals)

# 2) 载入主表，记录原列宽（按索引）
wb = load_workbook(TARGET)
ws = wb.active
old_w = {i: ws.column_dimensions[get_column_letter(i)].width for i in range(1, ws.max_column + 1)}

# 3) 插入空列
ws.insert_cols(INSERT_AT, 1)

# 4) 重排列宽：1..5 不变，6=新列，7.. = 原 6..
new_w = {}
for i in range(1, INSERT_AT):
    new_w[i] = old_w.get(i)
new_w[INSERT_AT] = 46
for i in range(INSERT_AT, ws.max_column):
    new_w[i + 1] = old_w.get(i)
for i, w in new_w.items():
    if w:
        ws.column_dimensions[get_column_letter(i)].width = w

# 5) 表头：复制相邻表头样式
hdr = ws.cell(1, INSERT_AT)
hdr.value = NEW_HEADER
hdr._style = copy.copy(ws.cell(1, INSERT_AT - 1)._style)
hdr.comment = Comment(
    "来源：中葡企业机构数据库_企业拜访表（列『出海诉求/跟进记录』），按公司名连接；同一公司多条拜访记录已去重合并。",
    "Claude",
)

# 6) 数据：按 公司简称 映射，并复制相邻列（本地化状态）单元格样式
matched, blank = [], []
for r in range(2, ws.max_row + 1):
    comp = ws.cell(r, 3).value  # 公司简称
    comp = str(comp).strip() if comp is not None else ""
    cell = ws.cell(r, INSERT_AT)
    cell._style = copy.copy(ws.cell(r, INSERT_AT - 1)._style)
    val = need.get(comp, "")
    cell.value = val if val else None
    (matched if val else blank).append(comp)

# 7) 扩展自动筛选范围（含新列）
last = get_column_letter(ws.max_column)
ws.auto_filter.ref = f"A1:{last}{ws.max_row}"

wb.save(TARGET)

uniq_matched = sorted(set(matched))
uniq_blank = sorted(set(blank))
print(f"已写入新列『{NEW_HEADER}』于第 {INSERT_AT} 列（{get_column_letter(INSERT_AT)}）")
print(f"行数: {ws.max_row-1}  有出海需求的行: {len(matched)}  空白行: {len(blank)}")
print(f"匹配到拜访记录的公司({len(uniq_matched)}): {uniq_matched}")
print(f"未匹配/留空的公司({len(uniq_blank)}): {uniq_blank}")
