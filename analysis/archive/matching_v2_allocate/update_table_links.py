# -*- coding: utf-8 -*-
"""用 PNCP 取回的链接更新三分类主表:补缺失的投标链接、新增『PNCP官方链接』(稳定可达)与『链接状态』。"""
import sys, io, json, re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.comments import Comment

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
F = "data/exports/中葡企业_巴西竞标匹配_三分类主表_v2.xlsx"
ref = json.load(open("analysis/matching/outputs/pncp_links_refetch.json", encoding="utf-8"))
conn = json.load(open("analysis/matching/outputs/link_connectivity.json", encoding="utf-8"))

def fix_scheme(u):
    if u is None: return None
    u = str(u).strip()
    if not u or u.lower() == "nan": return None
    if u.lower().startswith(("http://", "https://")): return u
    if "." in u.split("/")[0]:          # 缺协议但有域名 -> 补 https
        return "https://" + u
    return None                          # 纯相对路径,无法还原

def best_origin(pncp):
    r = ref.get(pncp, {})
    o = fix_scheme(r.get("orig_link"))
    return o or fix_scheme(r.get("linkSistemaOrigem"))

def status_of(pncp):
    r = ref.get(pncp, {}); c = conn.get(pncp, {})
    if not r.get("api_ok"):
        return "⚠ 已从PNCP下架/不可查(可能已撤销或结束)"
    if not best_origin(pncp):
        return "PNCP未登记源站链接 → 用PNCP官方链接"
    cls = c.get("origin_class", "")
    if "HTTP可达" in cls or "200" in cls: return "源站链接可访问"
    if "重定向" in cls:                    return "源站可访问(有跳转)"
    if "403" in cls:                       return "源站反爬:脚本被拒,浏览器可正常打开(或用PNCP官方链接)"
    if "失效" in cls:                      return "源站链接已失效 → 用PNCP官方链接"
    return "源站暂时无法连接 → 用PNCP官方链接"

wb = load_workbook(F); ws = wb.active
H = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
clink, cpncp = H["投标链接"], H["PNCP编号"]
c_portal = ws.max_column + 1
c_status = ws.max_column + 2

# 表头(套用既有深蓝白字样式)
hdr_fill = PatternFill("solid", fgColor="1F4E78"); hdr_font = Font(color="FFFFFF", bold=True, size=10)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
for ci, name in [(c_portal, "PNCP官方链接"), (c_status, "链接状态")]:
    cell = ws.cell(1, ci); cell.value = name
    cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = hdr_align
ws.cell(1, c_portal).comment = Comment(
    "PNCP 官方门户页面(pncp.gov.br),每个标的稳定权威入口,经测试全部可访问;"
    "源站(投标链接)打不开时,从此页可查看公告全文并跳转到原招标系统。", "Claude")
ws.cell(1, c_status).comment = Comment(
    "源站投标链接的连通性(自动检测于今日)。注:多数巴西采购平台为 SPA/需登录/反爬,"
    "'反爬'类链接在真人浏览器中可正常打开;'已下架/失效'请改用 PNCP官方链接。", "Claude")

filled = 0
for r in range(2, ws.max_row + 1):
    pncp = str(ws.cell(r, cpncp).value)
    rr = ref.get(pncp, {})
    # 补缺失投标链接
    cur = ws.cell(r, clink).value
    if (cur is None or not str(cur).strip() or str(cur).strip().lower() == "nan"):
        bo = best_origin(pncp)
        if bo:
            ws.cell(r, clink).value = bo; filled += 1
    # PNCP官方链接
    portal = rr.get("portal")
    ws.cell(r, c_portal).value = portal
    # 链接状态
    ws.cell(r, c_status).value = status_of(pncp)
    # 数据格样式
    for ci in (c_portal, c_status):
        ws.cell(r, ci).font = Font(size=9)
        ws.cell(r, ci).alignment = Alignment(vertical="top", wrap_text=True)

ws.column_dimensions[get_column_letter(c_portal)].width = 40
ws.column_dimensions[get_column_letter(c_status)].width = 30
ws.column_dimensions[get_column_letter(clink)].width = 40
ws.auto_filter.ref = f"A1:{get_column_letter(c_status)}{ws.max_row}"

wb.save(F)
print(f"已更新: {F}")
print(f"新增列: {get_column_letter(c_portal)}=PNCP官方链接, {get_column_letter(c_status)}=链接状态")
print(f"本次补上的投标链接(原为空): {filled} 行")
from collections import Counter
m = pd.read_excel(F, dtype=str)
print("投标链接 现非空:", m["投标链接"].notna().sum(), "/", len(m))
print("链接状态分布:")
for k, n in Counter(m["链接状态"]).most_common():
    print(f"  {n:3}  {k}")
