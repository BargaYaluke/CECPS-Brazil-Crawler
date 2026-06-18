# -*- coding: utf-8 -*-
"""ME/EPP 专属标处理:新增『ME/EPP限制』列标注全部行;剔除『大型企业 × 全部标项专属』行(移入单独sheet)。
   规模判定: capacity_tier>=3 视为大型(cap>=2000万,远超EPP R$4.8M营收上限)。
   法律注: 外资子公司/分支即便规模不大,据 LC123/2006 Art.3 §4(I/II) 通常也不具 ME/EPP 资格 -> 中型行加警示。"""
import sys, io, json
from copy import copy
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.comments import Comment

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
F = "data/exports/中葡企业_巴西竞标匹配_三分类主表_v2.xlsx"
me = json.load(open("analysis/matching/outputs/meepp_benefit.json", encoding="utf-8"))
comp = {c["公司简称"]: c for c in json.load(open("analysis/matching/outputs/company_master.json", encoding="utf-8"))}
prof = json.load(open("analysis/matching/outputs/company_profiles.json", encoding="utf-8"))

def tier_of(cn):
    c = comp.get(cn, {})
    return (prof.get(c.get("cid"), {}) or {}).get("capacity_tier")

def is_large(cn):
    t = tier_of(cn)
    return bool(t) and t >= 3   # tier3=cap2000万, tier4=cap5000万+ -> 大型,绝无 ME/EPP 资格

def meepp_label(pncp, cn):
    info = me.get(str(pncp), {})
    lab = info.get("label") or ""
    if "全部标项" in lab:
        if is_large(cn):
            return "仅限ME/EPP(全部标项)·大型企业不可投", "remove_large"
        return "⚠仅限ME/EPP(全部标项)·外资子公司通常不符资格,建议核实或改JV", "flag"
    if "部分标项" in lab:
        return lab.replace("ME/EPP专属", "仅限ME/EPP") + "·大企业仅可投其余标项", "flag"
    if "预留" in lab:
        return "含ME/EPP预留份额·大企业可投主份额", "ok"
    if "分包" in lab:
        return "要求向ME/EPP分包", "flag"
    return "", "none"

wb = load_workbook(F)
ws = wb.active
H = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
cpncp, ccomp = H["PNCP编号"], H["公司简称"]
cnew = ws.max_column + 1

# 表头(套既有深蓝白字)
hc = ws.cell(1, cnew); hc.value = "ME/EPP限制"
hc.fill = PatternFill("solid", fgColor="1F4E78"); hc.font = Font(color="FFFFFF", bold=True, size=10)
hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
hc.comment = Comment(
    "巴西 ME/EPP(中小企业)专属规则。'仅限ME/EPP(全部标项)'=整标只许巴西中小企业投,"
    "大型企业不可投;'部分标项'=仅部分标项保留给中小企业,大企业可投其余。"
    "注:据 LC123/2006 Art.3 §4,外资分支/有外资法人股东的子公司通常不具 ME/EPP 资格,"
    "故外资子公司即便规模不大也多半不能投全标专属标。", "Claude")

red = PatternFill("solid", fgColor="FFC7CE")
yel = PatternFill("solid", fgColor="FFF2CC")
remove_rows = []
for r in range(2, ws.max_row + 1):
    pncp = ws.cell(r, cpncp).value
    cn = ws.cell(r, ccomp).value
    label, action = meepp_label(pncp, cn)
    cell = ws.cell(r, cnew)
    cell.value = label
    cell.font = Font(size=9); cell.alignment = Alignment(vertical="top", wrap_text=True)
    if action == "remove_large":
        cell.fill = red; remove_rows.append(r)
    elif action == "flag":
        cell.fill = yel
ws.column_dimensions[get_column_letter(cnew)].width = 30

# 收集被剔除行(值)用于备份 sheet
ncol = ws.max_column
removed_vals = [[ws.cell(r, c).value for c in range(1, ncol + 1)] for r in remove_rows]
headers = [ws.cell(1, c).value for c in range(1, ncol + 1)]

# 自底向上删除
for r in sorted(remove_rows, reverse=True):
    ws.delete_rows(r, 1)
last = get_column_letter(ncol)
ws.auto_filter.ref = f"A1:{last}{ws.max_row}"

# 备份 sheet
bak = wb.create_sheet("已剔除_ME-EPP专属(大型企业)")
bak.append(headers)
for cell in bak[1]:
    cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.font = Font(color="FFFFFF", bold=True, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in removed_vals:
    bak.append(row)
for rr in range(2, bak.max_row + 1):
    for cc in range(1, ncol + 1):
        bak.cell(rr, cc).font = Font(size=9); bak.cell(rr, cc).alignment = Alignment(vertical="top", wrap_text=True)
for c in range(1, ncol + 1):
    bak.column_dimensions[get_column_letter(c)].width = ws.column_dimensions[get_column_letter(c)].width or 12
bak.freeze_panes = "A2"

wb.save(F)
print(f"已更新: {F}")
print(f"新增列 {get_column_letter(cnew)}=ME/EPP限制")
print(f"剔除(大型×全部专属) {len(remove_rows)} 行 -> 移入 sheet '已剔除_ME-EPP专属(大型企业)'")
m = pd.read_excel(F, sheet_name=0, dtype=str)
print(f"主表现 {len(m)} 行")
from collections import Counter
print("ME/EPP限制 分布(主表):")
for k, n in Counter(m["ME/EPP限制"].fillna("(无)")).most_common():
    print(f"  {n:3}  {k if k else '(无)'}")
print("\n剔除明细:")
for row in removed_vals:
    print(f"  {row[ccomp-1]:14} | {row[H['竞标中文梗概']-1][:30]} | {row[H['金额万CNY']-1]}万")
