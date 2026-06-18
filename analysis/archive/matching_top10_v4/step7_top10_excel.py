# -*- coding: utf-8 -*-
"""Step7(全覆盖+时效门槛版): 生成 中葡企业_巴西竞标匹配表_v4_时效筛选.xlsx。"""
import sys, io, json
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from match_lib import OUT, ROOT, TODAY

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

matches  = json.loads((OUT/"matches_top10.json").read_text("utf-8"))
overview = json.loads((OUT/"overview_top10.json").read_text("utf-8"))
audit    = json.loads((OUT/"tender_filter_audit.json").read_text("utf-8"))
tenders  = json.loads((OUT/"tenders_clean.json").read_text("utf-8"))
verdicts = json.loads((OUT/"verdicts.json").read_text("utf-8")) if (OUT/"verdicts.json").exists() else {}

df = pd.DataFrame(matches)
df["校验结论"] = df.apply(lambda r: verdicts.get(f"{r['cid']}|{r['PNCP编号']}",{}).get("结论",""), axis=1)
df["校验说明"] = df.apply(lambda r: verdicts.get(f"{r['cid']}|{r['PNCP编号']}",{}).get("说明",""), axis=1)
COLS = ["重要程度","公司简称","行业","本地化状态","所需本地化周期天","推荐优先级","可投性","时效可行",
        "首推","共享首推","竞标中文梗概","金额万CNY","采购方式","标的类型","截止日期","剩余天数",
        "州","级别","采购机构","价格登记SRP","长期挂网","契合分","可行性分","时效分","综合分",
        "时效渠道","落地建议","匹配理由","风险","校验结论","校验说明","PNCP编号","投标链接"]
df = df[[c for c in COLS if c in df.columns]]

# 精华可投表:全部企业里『时效可行 + 可投性=✓/○/△』的真正可推进项
ACTION = ["✓本轮可投","○时效可行(联合体/弱契合)","△可投但高风险(校验降级)"]
elite = df[(df["时效可行"]=="是") & (df["可投性"].isin(ACTION)) & (df["校验结论"]!="剔除")].copy()
elite = elite.sort_values(["重要程度","可投性","综合分"], ascending=[True, True, False])

ov = pd.DataFrame(overview)
ov = ov.sort_values(["imp_rank","本轮可投数"], ascending=[True, False]).drop(columns=["imp_rank"], errors="ignore")
OV_COLS = ["重要程度","公司简称","公司全称","行业","本地化状态","所需本地化周期天","政采契合","产能档位",
           "单标上限万元","本地化难度","本轮可投数","其中高风险","下一周期情报","校验剔除","低契合凑数",
           "首推标的","首推可投性","主营概括","落地提示"]
ov = ov[[c for c in OV_COLS if c in ov.columns]]

# 行业管线(含时效可行池大小)
tdf = pd.DataFrame(tenders)
import pandas as _pd
d = _pd.to_numeric(tdf["dias"], errors="coerce")
ADVICE = {
 "医疗医药":"药品/耗材/器械须ANVISA注册(高门槛,留足≥8个月);优先credenciamento常年挂网与年度SRP框架,本地化后入围滚动供货。",
 "高端制造":"设备须INMETRO等认证;本地组装/经销降税;主攻军方(EXERCITO/MARINHA)与州级设备credenciamento/SRP。",
 "大宗商贸":"建材/食品本地化PJ后挂网;以差异化进口品类切入,回避PNAE家庭农业专供标。",
 "数字经济":"多服务类;以软件许可/设备供货切入,关注联邦/州级IT设备credenciamento;纯本地运维需联合体。",
 "跨境商贸":"政采契合弱;以B2G货物经本地贸易商转供,提前布局PCA年度计划。",
 "文化体育":"仅器材供货可投(体育器材/乐器/灯光),活动服务类回避。",
}
pipe = pd.DataFrame([{
    "行业": dim, "可投标池": int((tdf["dim"]==dim).sum()),
    "时效充裕池(≥150天或无截止·货物)": int((((d>=150)|d.isna()) & (tdf["dim"]==dim) & (tdf["obj_type"]=="goods")).sum()),
    "credenciamento常年": int(((tdf["dim"]==dim)&(tdf["lt_bool"])).sum()),
    "未来管线建议": ADVICE[dim],
} for dim in ["医疗医药","高端制造","大宗商贸","数字经济","跨境商贸","文化体育"]])

method = pd.DataFrame([
 ("交付口径","全部437家企业(438登记行,上海擎朗智能重复1行)每家恰好10个竞标;按重要程度排序"),
 ("时效基准日", TODAY + " (据此倒算各企业可否赶上投标)"),
 ("本地化时效门槛(核心)","已在巴西14天/葡语区90/需本地化:低90·中150·高240天。标的『剩余天数 < 该周期』即判定『截止过短·本轮不现实』,降级为下一周期需求情报"),
 ("数据来源","PNCP招标 report.xlsx(8357,采集2026-05-22~28)+中葡企业机构数据库(437家)"),
 ("评分引擎","DeepSeek deepseek-v4-flash(企业画像+逐标契合/可行/时效打分)"),
 ("过滤口径"," | ".join(f"{k}:{v}" for k,v in audit.items())),
 ("综合分","0.45契合+0.30可行+0.25时效;同企业先按可投性档位、再综合分排序取top10"),
 ("可投性分层","✓本轮可投 / ○时效可行(联合体·弱) / △可投但高风险(校验降级) / ◇下一周期·需求情报(截止过短,产品对口可作下一轮目标) / ·低契合(凑数,为凑满10个的弱相关) / ✕不建议(校验剔除)"),
 ("关键结论","本快照近端标为主(中位剩12天),时效充裕(≥150天/无截止)货物标全网仅~179个;故多数需本地化企业本轮『可投』极少甚至为0,首要任务是本地化以捕获下一周期"),
 ("对抗校验","对高/中重要度企业 top10 中『时效可行+契合≥5』的本轮可投核心,做多智能体(Claude)独立校验;剔除/降级已标色"),
 ("用法","先看『本轮可投精华』表=可立即推进项;总表按『时效可行=是』筛即各企业现实可投;◇情报用于排本地化优先级;金额含数据坑,以PNCP原始Edital为准"),
], columns=["维度","说明"])

out_path = ROOT/"data/exports/中葡企业_巴西竞标匹配表_v4_时效筛选.xlsx"
with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
    elite.to_excel(xw, sheet_name="本轮可投精华(时效可行)", index=False)
    df.to_excel(xw, sheet_name="竞标匹配总表(每企业10标)", index=False)
    ov.to_excel(xw, sheet_name="企业匹配总览", index=False)
    pipe.to_excel(xw, sheet_name="行业未来管线建议", index=False)
    method.to_excel(xw, sheet_name="方法论与口径", index=False)

    wb = xw.book
    hdr_fill = PatternFill("solid", fgColor="1F4E78"); hdr_font = Font(color="FFFFFF", bold=True, size=10)
    imp_fill = {"1高":"FFC7CE","2中":"FFEB9C","3低":"D9E1F2","无标识(最低)":"F2F2F2"}
    tier_fill = {"✓本轮可投":"C6EFCE","○时效可行(联合体/弱契合)":"DDEBF7","△可投但高风险(校验降级)":"FCE4D6",
                 "◇下一周期·需求情报(截止过短)":"FFF2CC","·低契合(凑数)":"FFFFFF","✕不建议(校验剔除)":"D9D9D9"}
    feas_fill = {"是":"C6EFCE","否(截止过短)":"FCE4D6"}
    vd_fill = {"剔除":"FFC7CE","降级":"FFEB9C","通过":"C6EFCE"}
    WIDTHS = {"竞标中文梗概":46,"匹配理由":36,"落地建议":26,"风险":22,"采购机构":24,"采购方式":17,
              "投标链接":32,"PNCP编号":24,"主营概括":22,"落地提示":24,"未来管线建议":56,"说明":92,
              "首推标的":34,"可投性":18,"时效可行":13,"所需本地化周期天":10,"时效渠道":14,"首推可投性":16}
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = hdr_fill; cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hdrs = {ws.cell(1,ci).value: ci for ci in range(1, ws.max_column+1)}
        for ci in range(1, ws.max_column+1):
            ws.column_dimensions[get_column_letter(ci)].width = WIDTHS.get(ws.cell(1,ci).value, 12)
        for ri in range(2, ws.max_row+1):
            for ci in range(1, ws.max_column+1):
                ws.cell(ri,ci).alignment = Alignment(vertical="top", wrap_text=True)
                ws.cell(ri,ci).font = Font(size=9)
            def paint(col, mp):
                if col in hdrs:
                    v = str(ws.cell(ri,hdrs[col]).value or "").strip()
                    if v in mp and mp[v] != "FFFFFF":
                        ws.cell(ri,hdrs[col]).fill = PatternFill("solid", fgColor=mp[v])
            paint("重要程度", imp_fill); paint("可投性", tier_fill)
            paint("时效可行", feas_fill); paint("校验结论", vd_fill)
            if "首推" in hdrs and str(ws.cell(ri,hdrs["首推"]).value or "").strip():
                ws.cell(ri,hdrs["首推"]).font = Font(size=9, bold=True, color="BF8F00")

import collections
print("已生成:", out_path)
print("总表行:", len(df), "| 精华可投行:", len(elite), "| 企业:", len(ov))
print("可投性分布:", dict(collections.Counter(df["可投性"])))
print("精华可投·各重要程度:", dict(collections.Counter(elite["重要程度"])))
