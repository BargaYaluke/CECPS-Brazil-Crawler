# -*- coding: utf-8 -*-
"""Step9: 生成单一主表(三分类) —— 中葡企业_巴西竞标匹配_三分类主表.xlsx。
   仅保留重要特征:删除 首推/共享首推/级别/各类分数。"""
import sys, io, json, glob, os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from match_lib import OUT, ROOT

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

comp = {c["cid"]: c for c in json.loads((OUT/"company_master.json").read_text("utf-8"))}
cat1 = json.loads((OUT/"cat1.json").read_text("utf-8"))
cat2 = json.loads((OUT/"cat2.json").read_text("utf-8"))
cat3 = json.loads((OUT/"cat3.json").read_text("utf-8"))

# 合并校验裁决
v1, v2 = {}, {}
for f in glob.glob(str(OUT/"cat_verdicts"/"*.json")):
    try: d = json.loads(open(f, encoding="utf-8").read())
    except Exception: continue
    cid = d.get("cid") or os.path.splitext(os.path.basename(f))[0]
    for x in d.get("cat1", []):
        if x.get("pncp"): v1[f"{cid}|{x['pncp']}"] = x
    for x in d.get("cat2", []):
        if x.get("pncp"): v2[f"{cid}|{x['pncp']}"] = x

def amt(r): return r.get("金额万")
def imp(cid): return comp[cid]["重要程度"]
def imprank(cid): return comp[cid]["imp_rank"]

rows = []
# 第一类:核心匹配(去掉被校验剔除的)
kept1 = drop1 = 0
for r in cat1:
    vd = v1.get(f"{r['cid']}|{r['pncp']}")
    if vd and vd.get("结论") == "剔除":
        drop1 += 1; continue
    kept1 += 1
    rows.append({
        "类别": "① 核心匹配(可投)", "_sort": (1, imprank(r["cid"])),
        "重要程度": imp(r["cid"]), "公司简称": comp[r["cid"]]["公司简称"],
        "行业": comp[r["cid"]].get("行业") or "(未分类)", "本地化状态": comp[r["cid"]]["本地化状态"],
        "标的类型": r.get("类型"), "竞标中文梗概": r.get("梗概"), "金额万CNY": amt(r),
        "采购方式": r.get("采购方式"), "剩余天数": ("无截止" if r.get("dias") is None else r.get("dias")),
        "州": r.get("州"), "价格登记SRP": "是" if r.get("srp") else "否",
        "长期挂网": "是" if r.get("lt") else "否",
        "落地建议": r.get("channel"), "匹配理由": r.get("reason"),
        "采购机构": r.get("机构"), "PNCP编号": r.get("pncp"), "投标链接": r.get("link"),
    })
# 第二类:本地化+牌照后可争取(只留 可争取)
kept2 = drop2 = 0
for r in cat2:
    vd = v2.get(f"{r['cid']}|{r['pncp']}")
    if vd and vd.get("结论") == "不可行":
        drop2 += 1; continue
    kept2 += 1
    zizhi = (vd or {}).get("所需资质") or "相应本地牌照/资质"
    note = (vd or {}).get("说明") or r.get("reason")
    rows.append({
        "类别": "② 本地化+牌照后可争取", "_sort": (2, imprank(r["cid"])),
        "重要程度": imp(r["cid"]), "公司简称": comp[r["cid"]]["公司简称"],
        "行业": comp[r["cid"]].get("行业") or "(未分类)", "本地化状态": comp[r["cid"]]["本地化状态"],
        "标的类型": r.get("类型"), "竞标中文梗概": r.get("梗概"), "金额万CNY": amt(r),
        "采购方式": r.get("采购方式"), "剩余天数": ("无截止" if r.get("dias") is None else r.get("dias")),
        "州": r.get("州"), "价格登记SRP": "是" if r.get("srp") else "否",
        "长期挂网": "是" if r.get("lt") else "否",
        "落地建议": f"本地化+取得【{zizhi}】后争取", "匹配理由": note,
        "采购机构": r.get("机构"), "PNCP编号": r.get("pncp"), "投标链接": r.get("link"),
    })
# 第三类:高价值候选(不绑定企业)
for r in cat3:
    rows.append({
        "类别": "③ 高价值候选(待JV/本地伙伴)", "_sort": (3, -(r.get("金额万") or 0)),
        "重要程度": "", "公司简称": "—(高价值候选,不限企业)",
        "行业": r.get("行业"), "本地化状态": "",
        "标的类型": "", "竞标中文梗概": r.get("梗概"), "金额万CNY": r.get("金额万"),
        "采购方式": r.get("采购方式"), "剩余天数": ("无截止" if r.get("dias") is None else r.get("dias")),
        "州": r.get("州"), "价格登记SRP": "是" if r.get("srp") else "否",
        "长期挂网": "是" if r.get("lt") else "否",
        "落地建议": "高价值大单:建议组联合体(consórcio)/找本地伙伴或本地化后争取",
        "匹配理由": "金额高+截止充裕,值得跟踪", "采购机构": r.get("机构"),
        "PNCP编号": r.get("pncp"), "投标链接": r.get("link"),
    })

df = pd.DataFrame(rows).sort_values(["_sort","公司简称"]).drop(columns=["_sort"])
COLS = ["类别","重要程度","公司简称","行业","本地化状态","标的类型","竞标中文梗概","金额万CNY",
        "采购方式","剩余天数","州","价格登记SRP","长期挂网","落地建议","匹配理由","采购机构","PNCP编号","投标链接"]
df = df[COLS]

out_path = ROOT/"data/exports/中葡企业_巴西竞标匹配_三分类主表.xlsx"
with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
    df.to_excel(xw, sheet_name="竞标匹配主表(三分类)", index=False)
    ws = xw.book.active
    hdr_fill = PatternFill("solid", fgColor="1F4E78"); hdr_font = Font(color="FFFFFF", bold=True, size=10)
    cat_fill = {"① 核心匹配(可投)":"C6EFCE", "② 本地化+牌照后可争取":"FFF2CC", "③ 高价值候选(待JV/本地伙伴)":"DDEBF7"}
    imp_fill = {"1高":"FFC7CE","2中":"FFEB9C","3低":"D9E1F2"}
    WIDTHS = {"类别":20,"竞标中文梗概":50,"匹配理由":34,"落地建议":30,"采购机构":24,"采购方式":17,
              "投标链接":32,"PNCP编号":24,"公司简称":16,"行业":10,"本地化状态":10}
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hdrs = {ws.cell(1,ci).value: ci for ci in range(1, ws.max_column+1)}
    for ci in range(1, ws.max_column+1):
        ws.column_dimensions[get_column_letter(ci)].width = WIDTHS.get(ws.cell(1,ci).value, 12)
    for ri in range(2, ws.max_row+1):
        for ci in range(1, ws.max_column+1):
            ws.cell(ri,ci).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(ri,ci).font = Font(size=9)
        cv = str(ws.cell(ri, hdrs["类别"]).value or "")
        if cv in cat_fill:
            ws.cell(ri, hdrs["类别"]).fill = PatternFill("solid", fgColor=cat_fill[cv])
        iv = str(ws.cell(ri, hdrs["重要程度"]).value or "")
        if iv in imp_fill:
            ws.cell(ri, hdrs["重要程度"]).fill = PatternFill("solid", fgColor=imp_fill[iv])

from collections import Counter
print("已生成:", out_path)
print("主表行数:", len(df))
print("分类分布:", dict(Counter(df["类别"])))
print(f"第一类 校验保留{kept1}/剔除{drop1} | 第二类 保留{kept2}/剔除{drop2}")
