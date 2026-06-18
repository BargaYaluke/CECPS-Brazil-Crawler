# -*- coding: utf-8 -*-
"""Step5: 生成『巴西金主(KA)竞标分析表』。
主表(1519标):一级/二级/需求类型/长期挂网 + 标的信息 + 每标 top5 强出海中企(经Claude核验)。
金主清单(200家):KA画像。
每标 top5 = 从该桶经核验名册中,按 需求类型↔能力 过滤,再按 出海强度→巴西→体量契合 排序取5。
"""
import sys, io, json, math
from collections import Counter
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OUT = HERE / "outputs"

bids = json.loads((OUT/"ka_bids.json").read_text("utf-8"))
cls = json.loads((OUT/"classifications.json").read_text("utf-8"))
bucket_index = json.loads((OUT/"bucket_index.json").read_text("utf-8"))
rosters = json.loads((OUT/"rosters.json").read_text("utf-8"))
buyers = pd.read_csv(OUT/"ka_buyers.csv")
amt_audit = json.loads((OUT/"amount_audit.json").read_text("utf-8"))
amt_cat = {r["pncp"]: r.get("verify_cat", "") for r in amt_audit["rows"]}

CAP_FOR_DEMAND = {
    "产品采购": {"产品", "综合"},
    "工程总包": {"工程", "综合"},
    "服务外包": {"服务", "运营", "综合"},
    "租赁运营": {"运营", "产品", "综合"},
    "维修维护": {"服务", "运营", "产品", "综合"},
    "系统建设": {"系统", "综合", "产品"},
    "长期入库/征召": {"产品", "服务", "综合"},
    "其他": None,
}
SIZE_RANK = {"大": 3, "中": 2, "小": 1}


def bid_scale(valor_wan):
    if not valor_wan:
        return "small"
    if valor_wan >= 2000:
        return "big"
    if valor_wan >= 300:
        return "mid"
    return "small"


def size_fit(scale, size):
    """体量契合软分(大单偏大型;小单不强求大型)。"""
    s = SIZE_RANK.get(size, 2)
    if scale == "big":
        return s            # 3>2>1, 偏大
    if scale == "mid":
        return 3 - abs(s - 2.5)  # 中性偏中大
    return 4 - s            # small: 小/中略优先,大仍可


def pick_top5(bid):
    bkey = "|".join(bucket_index.get(bid["pncp"], ["其他", "综合"]))
    roster = rosters.get(bkey, [])
    if not roster:
        return [], bkey
    demand = cls.get(bid["pncp"], {}).get("demand", "其他")
    allow = CAP_FOR_DEMAND.get(demand)
    scale = bid_scale(bid.get("valor_wan"))
    prim = [r for r in roster if (allow is None or r["cap"] in allow)]
    rest = [r for r in roster if r not in prim]
    def score(r):
        return (r.get("ovs_rank", 0), r.get("brazil", False), size_fit(scale, r.get("size", "中")))
    prim.sort(key=lambda r: score(r), reverse=True)
    rest.sort(key=lambda r: score(r), reverse=True)
    chosen = (prim + rest)[:5]
    return chosen, bkey


def comp_line(i, r):
    tags = [r.get("size", "中"), "出海" + r.get("overseas", "未知")]
    if r.get("brazil"):
        tags.append("巴西✓")
    if not r.get("verified", True):
        tags.append("未核验")
    ev = r.get("evidence") or r.get("fit") or "—"
    return f"{['①','②','③','④','⑤'][i]} {r['name']}({'·'.join(tags)})— {ev}"


rows = []
no_roster = 0
for b in bids:
    top5, bkey = pick_top5(b)
    if not top5:
        no_roster += 1
        rec = "(该细分暂无经核验的强出海中企 → 建议人工拓展或组联合体)"
    else:
        rec = "\n".join(comp_line(i, r) for i, r in enumerate(top5))
    c = cls.get(b["pncp"], {})
    rows.append({
        "KA排名": b["ka_rank"],
        "采购机构(金主)": b["orgao"],
        "级别": b["esfera"],
        "州": b["uf"],
        "城市": b.get("municipio"),
        "一级行业": c.get("l1", "其他"),
        "二级专业领域": c.get("l2", "未细分"),
        "需求类型": c.get("demand", "其他"),
        "长期挂网": b["long_term"],
        "价格登记SRP": b["srp"],
        "竞标中文梗概": b["obj_zh"],
        "金额万CNY": ("预算保密" if b.get("sigilo") else b.get("valor_wan")),
        "金额核验": amt_cat.get(b["pncp"], ""),
        "提交截止日期": b["deadline"],
        "剩余天数": ("无截止" if b.get("dias") is None else b["dias"]),
        "ME/EPP限制": b["meepp"],
        "推荐中国企业TOP5(强出海意愿·经核验)": rec,
        "PNCP编号": b["pncp"],
        "投标链接": b.get("link"),
        "PNCP官方链接": b.get("portal"),
    })

df = pd.DataFrame(rows)
# 排序:金额从大到小;预算保密/象征性(无真实金额)沉底
df["_amt"] = pd.to_numeric(df["金额万CNY"], errors="coerce")
df["_bottom"] = df["金额核验"].isin(["预算保密(sigilo)", "象征性/未计价"]).astype(int)
df = df.sort_values(["_bottom", "_amt", "KA排名"],
                    ascending=[True, False, True],
                    na_position="last").drop(columns=["_amt", "_bottom"]).reset_index(drop=True)

COLS = list(df.columns)
REC_COL = "推荐中国企业TOP5(强出海意愿·经核验)"

# ── 金主清单 sheet ──
buyers = buyers.rename(columns={
    "ka_rank": "KA排名", "采购机构": "采购机构(金主)", "uf": "州", "gdp": "GDP分层",
    "n_bids": "发标数(可投池)", "total_wan": "总金额万CNY", "median_wan": "单标中位万CNY",
    "max_wan": "单标最大万CNY", "ka_score": "KA综合分"})
# 标数(本表纳入的未过期标)
inc = Counter(b["orgao"] for b in bids)
buyers["本表纳入标数"] = buyers["采购机构(金主)"].map(lambda x: inc.get(x, 0))
buyers = buyers[["KA排名", "采购机构(金主)", "级别", "州", "GDP分层", "发标数(可投池)",
                 "本表纳入标数", "总金额万CNY", "单标中位万CNY", "单标最大万CNY", "KA综合分"]]

# ── 写 Excel + 样式 ──
WIDTHS = {"KA排名": 7, "采购机构(金主)": 30, "级别": 7, "州": 6, "城市": 14,
          "一级行业": 12, "二级专业领域": 14, "需求类型": 13, "长期挂网": 9, "价格登记SRP": 11,
          "竞标中文梗概": 52, "金额万CNY": 12, "金额核验": 20, "提交截止日期": 17, "剩余天数": 9, "ME/EPP限制": 30,
          REC_COL: 75, "PNCP编号": 25, "投标链接": 38, "PNCP官方链接": 44}
HDR_COMMENTS = {
    "采购机构(金主)": "金主=综合分(发标频次×总金额×单标中位额)排名 top-200 的采购机构;详见『金主KA清单』sheet。",
    "一级行业": "据标的梗概+葡语原文 由 DeepSeek 分类,固定一级行业表择一。",
    "二级专业领域": "细分到可指导找供应商的专业领域(如 医疗设备/超声诊断/义齿制作/道路工程/校餐食品)。",
    "需求类型": "履约模式:产品采购/服务外包/工程总包/租赁运营/维修维护/系统建设/长期入库·征召。不同类型投入与适配差异大。",
    "金额核验": "金额已交叉核验(report.xlsx 预估金额 = DB 抓取值,完全一致;FX 汇率统一 1 BRL≈1.341 CNY)。分类:『逐项预算一致✓』=金额与逐项预算汇总吻合,硬值可信;『估算上限·逐项预算仅部分』=价格登记(SRP)/多年期合同的框架估算上限,我方抓取的逐项预算仅覆盖部分,金额为官方申报的年度/合同上限(非笔误);『预算保密』=机构未披露(sigilo),金额=0;『象征性/未计价』=credenciamento 占位值(≈0)。无发现笔误虚高。",
    "长期挂网": "是=credenciamento 公开征召/长期机会标,随时可投、一次入围长期复购,KA 经营首选。",
    "价格登记SRP": "是=走价格登记制(registro de preços):一次招标登记入库,有效期内按需分批下单,其他机构可搭车(carona)。",
    "提交截止日期": "源表真实投标截止(巴西利亚时间);『无截止』为长期挂网/价格登记类。仅含未过期标。",
    "ME/EPP限制": "巴西中小企业专属规则。『全部专属』=整标只许巴西中小企业投,外资子公司通常不符资格(LC123/2006 Art.3§4),建议核实或改JV;『部分专属/预留』大企业可投其余/主份额。",
    REC_COL: "每标推荐5家最匹配的中国企业(不限于现有名录,全中国范围、优先出海意愿强者)。流程:DeepSeek 按细分领域海选 → Claude 逐家联网核验(真实存在+出海动作),按 出海强度→巴西布局→体量契合 排序,并按本标『需求类型』过滤履约能力。标注:体量(大/中/小)·出海强度(强=已在巴西拉美有实体项目/有=有出口或海外项目/弱=初步出海)·巴西✓=有巴西或拉美布局。出海证据/适配附后。『未核验』=核验环节未覆盖,供参考。",
    "PNCP官方链接": "PNCP 官方门户(pncp.gov.br),每标稳定权威入口,可看公告全文并跳转原系统。",
    "投标链接": "原招标系统链接(DB采集);部分为 SPA/需登录/反爬,真人浏览器可开;打不开时改用 PNCP官方链接。",
}
hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(color="FFFFFF", bold=True, size=10)
hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
l1_fills = {}  # 一级行业 配色
PALETTE = ["FCE4D6", "DDEBF7", "E2EFDA", "FFF2CC", "EDEDED", "FBE2D5", "D6DCE4", "E1D5E7",
           "D5E8D4", "FFE6CC", "DAE8FC", "F8CECC", "D5F5E3", "FDEBD0", "EBDEF0", "D4E6F1",
           "FADBD8", "E8DAEF", "D1F2EB", "FCF3CF", "F6DDCC"]
yel = PatternFill("solid", fgColor="FFF2CC"); red = PatternFill("solid", fgColor="FFC7CE")
grn = PatternFill("solid", fgColor="C6EFCE")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_main(ws, n_rows, hdrs):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n_rows+1}"
    for cell in ws[1]:
        cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = hdr_align; cell.border = border
    for name, txt in HDR_COMMENTS.items():
        if name in hdrs:
            ws.cell(1, hdrs[name]).comment = Comment(txt, "Claude")
    for ci in range(1, len(COLS)+1):
        ws.column_dimensions[get_column_letter(ci)].width = WIDTHS.get(ws.cell(1, ci).value, 12)
    l1s = sorted({ws.cell(r, hdrs["一级行业"]).value for r in range(2, n_rows+2)})
    for i, l1 in enumerate(l1s):
        l1_fills[l1] = PatternFill("solid", fgColor=PALETTE[i % len(PALETTE)])
    for ri in range(2, n_rows+2):
        for ci in range(1, len(COLS)+1):
            cell = ws.cell(ri, ci)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(size=9); cell.border = border
        ws.cell(ri, hdrs["一级行业"]).fill = l1_fills.get(ws.cell(ri, hdrs["一级行业"]).value)
        mv = str(ws.cell(ri, hdrs["ME/EPP限制"]).value or "")
        if "全部专属" in mv:
            ws.cell(ri, hdrs["ME/EPP限制"]).fill = red
        elif mv and mv != "无":
            ws.cell(ri, hdrs["ME/EPP限制"]).fill = yel
        if str(ws.cell(ri, hdrs["长期挂网"]).value) == "是":
            ws.cell(ri, hdrs["长期挂网"]).fill = grn
        if "金额核验" in hdrs:
            vv = str(ws.cell(ri, hdrs["金额核验"]).value or "")
            if vv.startswith("估算上限"):
                ws.cell(ri, hdrs["金额核验"]).fill = yel
            elif vv.startswith("预算保密") or vv.startswith("象征性"):
                ws.cell(ri, hdrs["金额核验"]).fill = PatternFill("solid", fgColor="D9D9D9")


def style_buyers(ws, n_rows, ncol):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{n_rows+1}"
    for cell in ws[1]:
        cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = hdr_align
    bw = {1: 7, 2: 34, 3: 7, 4: 6, 5: 9, 6: 13, 7: 12, 8: 14, 9: 14, 10: 14, 11: 10}
    for ci in range(1, ncol+1):
        ws.column_dimensions[get_column_letter(ci)].width = bw.get(ci, 12)
    for ri in range(2, n_rows+2):
        for ci in range(1, ncol+1):
            ws.cell(ri, ci).font = Font(size=9)
            ws.cell(ri, ci).alignment = Alignment(vertical="center")


out_path = ROOT/"data/exports/巴西金主KA竞标分析表.xlsx"
with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
    df.to_excel(xw, sheet_name="金主竞标主表", index=False)
    ws = xw.book["金主竞标主表"]
    hdrs = {ws.cell(1, ci).value: ci for ci in range(1, ws.max_column+1)}
    style_main(ws, len(df), hdrs)
    buyers.to_excel(xw, sheet_name="金主KA清单(200家)", index=False)
    ws2 = xw.book["金主KA清单(200家)"]
    style_buyers(ws2, len(buyers), buyers.shape[1])

print("已生成:", out_path)
print("主表行数:", len(df), "| 无名册标:", no_roster)
print("一级行业分布:", dict(Counter(df["一级行业"]).most_common()))
print("需求类型分布:", dict(Counter(df["需求类型"]).most_common()))
print("长期挂网=是:", int((df["长期挂网"] == "是").sum()), "| SRP=是:", int((df["价格登记SRP"] == "是").sum()))
print("ME/EPP 全部专属:", int((df["ME/EPP限制"].str.contains("全部专属")).sum()))
