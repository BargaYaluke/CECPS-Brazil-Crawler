# -*- coding: utf-8 -*-
"""
设备/产品采购标筛选(只用标的本身信息,不做企业匹配)

口径:
  1) 未过期      data_encerramento_proposta > 今天(2026-06-18)
  2) 金额区间    20万 ≤ valor_cny_estimado ≤ 200万 CNY(1 BRL≈1.34135 CNY)
  3) 实物采购    itens 中 Material(M)金额占比 ≥ 0.8  →  候选池
                 ⚠ 源数据部分服务项被误标 M(如理疗/住宿/仲裁/铺路),
                 故再用 DeepSeek 按 objeto+物品描述判定『是否实物采购』并归类,
                 剔除服务/工程施工,得最终设备/产品标。
  4) 剔拍卖      Leilão 系政府资产变卖,非采购,排除

产物:data/exports/巴西设备产品采购标_20-200万CNY.xlsx
"""
import sys, json, asyncio
from pathlib import Path
from datetime import date
import sqlite3
import pandas as pd
import httpx
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\巴西爬虫")
sys.path.insert(0, str(ROOT / "analysis" / "matching" / "pipeline"))
import match_lib as ml  # 复用 DeepSeek 客户端 + 磁盘缓存

# 模式:默认 active(仅未过期);传 all → 不限时效(含已过期)
MODE = "all" if len(sys.argv) > 1 and sys.argv[1].lower() in ("all", "全部", "不限") else "active"

DB = ROOT / "data" / "procurement.db"
_suffix = "_全量含过期" if MODE == "all" else ""
OUT = ROOT / "data" / "exports" / f"巴西设备产品采购标_20-200万CNY{_suffix}.xlsx"
CLS_JSON = ROOT / "analysis" / "equipment_bids" / f"classifications{_suffix}.json"
TODAY = "2026-06-18"
CNY_LO, CNY_HI = 200_000, 2_000_000
MAT_SHARE = 0.8

CATS = ["车辆/工程机械", "工业/生产设备", "IT/电子/通信设备", "医疗器械/耗材",
        "实验/检测仪器", "办公/教学/家具", "厨房/家电", "建材/五金/物料",
        "药品", "食品/农产品", "燃料/化工品", "其他实物",
        "服务(非实物)", "工程施工"]
NON_PRODUCT = {"服务(非实物)", "工程施工"}

# ── 1) 候选池:金额区间 + Material≥0.8,剔拍卖;active 模式再加未过期 ─────
_active_clause = f"AND data_encerramento_proposta > '{TODAY}'" if MODE == "active" else ""
con = sqlite3.connect(DB)
c = pd.read_sql_query(f"""
  SELECT pncp_id, objeto_compra, valor_total_estimado, valor_cny_estimado,
         data_publicacao_pncp, data_encerramento_proposta, modalidade_nome,
         srp, is_long_term_opportunity, orgao_razao_social, orgao_esfera_id,
         uf_sigla, municipio_nome, link_sistema_origem
  FROM contratacoes
  WHERE valor_cny_estimado BETWEEN {CNY_LO} AND {CNY_HI}
    {_active_clause}
    AND modalidade_nome NOT LIKE '%Leil%'
""", con)
print(f"[模式 {MODE}] 金额区间内标 {len(c)} 条")
it = pd.read_sql_query("""SELECT pncp_id, descricao, quantidade, unidade_medida,
       valor_total, material_ou_servico, tipo_beneficio_nome FROM itens""", con)
tc = pd.read_sql_query("SELECT source_text, translated FROM translation_cache", con)
con.close()

it = it[it.pncp_id.isin(set(c.pncp_id))].copy()
it["valor_total"] = it["valor_total"].fillna(0)
share = (it.groupby("pncp_id")
         .apply(lambda g: g.loc[g.material_ou_servico == "M", "valor_total"].sum()
                / (g["valor_total"].sum() or 1)).rename("_mat_share"))
c = c.merge(share, left_on="pncp_id", right_index=True, how="left")
c["_mat_share"] = c["_mat_share"].fillna(0)
c = c[c["_mat_share"] >= MAT_SHARE].reset_index(drop=True)

# 主要采购物品(Material 项,按金额前 4)+ 给 LLM 的物品描述串
def item_lines(pncp, n, with_qty):
    g = it[(it.pncp_id == pncp) & (it.material_ou_servico == "M")].sort_values("valor_total", ascending=False)
    out = []
    for _, r in g.head(n).iterrows():
        d = str(r.descricao or "").strip().replace("\n", " ").replace("\r", " ")[:70]
        if with_qty:
            q = "" if pd.isna(r.quantidade) else f"({r.quantidade:g}{r.unidade_medida or ''})"
            out.append(d + q)
        else:
            out.append(d)
    return " | ".join(out)
c["主要采购物品"] = [item_lines(p, 4, True) for p in c.pncp_id]
c["_items_pt"] = [item_lines(p, 6, False) for p in c.pncp_id]

# ── 2) DeepSeek 归类:是否实物采购 + 类别 ───────────────────────────────
SYS = (
    "你是巴西政府采购分类助手。给你若干招标(objeto葡萄牙语标题 + 主要标的物品葡语描述),"
    "判断每条的【真实采购性质】并归类。注意:源数据可能把服务误标成物品,以文本语义为准。\n"
    "类别(cat 只能取其一):" + "、".join(CATS) + "。\n"
    "判定 is_prod=true 当且仅当该标【主体是采购实物货物】(设备/机械/车辆/器械/IT硬件/家具/"
    "家电/建材物料/药品/食品/燃料等可交付的实物商品);\n"
    "is_prod=false 当主体是【服务】(如理疗、住宿餐饮、仲裁裁判、运输、维护保养、咨询、人力、"
    "保险、认证credenciamento提供服务等)或【工程施工obras】(道路铺装recapeamento、建筑construção、"
    "改造reforma、管沟安装施工等)。\n"
    "只输出 JSON: {\"items\":[{\"idx\":int,\"cat\":\"类别\",\"is_prod\":bool}, ...]} ,每个 idx 必须出现一次。"
)

def build_user(batch):
    lines = []
    for idx, obj, items in batch:
        obj = (obj or "").replace("\n", " ").replace("\r", " ")[:240]
        items = (items or "")[:360]
        lines.append(f'idx={idx}\nobjeto: {obj}\nitens: {items}')
    return "请分类以下招标:\n\n" + "\n---\n".join(lines)

async def classify_all(records):
    BATCH = 12
    batches = [records[i:i + BATCH] for i in range(0, len(records), BATCH)]
    async with httpx.AsyncClient() as client:
        factories = [
            (lambda s, b=b: ml.deepseek_json(client, SYS, build_user(b), sem=s))
            for b in batches
        ]
        results = await ml.run_pool(factories, concurrency=6)
    out = {}
    for res in results:
        if not res:
            continue
        for o in (res.get("items") if isinstance(res, dict) else res) or []:
            try:
                out[int(o["idx"])] = {"cat": o.get("cat"), "is_prod": bool(o.get("is_prod"))}
            except Exception:
                pass
    return out

records = [(i, r.objeto_compra, r._items_pt) for i, r in c.iterrows()]
cls = asyncio.run(classify_all(records))
CLS_JSON.write_text(json.dumps(cls, ensure_ascii=False, indent=1), encoding="utf-8")

# 关键词兜底(LLM 未覆盖到的 idx)
def kw_cat(o):
    t = (o or "").lower()
    if any(k in t for k in ["aliment", "merenda", "hortifrut", "frutas", "verdura", "gênero alim", "carne", "leite", "panific", "agricultura familiar", "perecív"]): return "食品/农产品"
    if any(k in t for k in ["medicament", "farmac"]): return "药品"
    if any(k in t for k in ["hospitalar", "odontolog", "órtese", "prótese", "cirúrg", "insumo", "seringa", "sonda"]): return "医疗器械/耗材"
    if any(k in t for k in ["veícul", "veicul", "ônibus", "caminh", "ambulânc", "automot", "trator", "retroescav"]): return "车辆/工程机械"
    if any(k in t for k in ["informátic", "computador", "periféric", "notebook", "servidor", "eletrôni"]): return "IT/电子/通信设备"
    if any(k in t for k in ["equipament", "máquina", "mobiliári", "móveis", "utensíli", "eletrodom", "climatiz"]): return "工业/生产设备"
    if any(k in t for k in ["construção", "obra", "pavimenta", "reforma", "engenharia", "recapeament", "asfalt"]): return "工程施工"
    if any(k in t for k in ["serviç", "fisioterap", "arbitrag", "diária", "hospedagem", "locação", "manutenção", "consultoria"]): return "服务(非实物)"
    return "其他实物"

c["物资类别"] = [
    (cls.get(i, {}).get("cat") or kw_cat(r.objeto_compra)) for i, r in c.iterrows()
]
c["_is_prod"] = [
    cls[i]["is_prod"] if i in cls else (kw_cat(r.objeto_compra) not in NON_PRODUCT)
    for i, r in c.iterrows()
]
# 类别非法值回退
c.loc[~c["物资类别"].isin(CATS), "物资类别"] = [kw_cat(o) for o in c.loc[~c["物资类别"].isin(CATS), "objeto_compra"]]

# ── 3) 拆分:产品标(主表) / 疑似服务工程(副表) ─────────────────────
prod = c[c["_is_prod"] & ~c["物资类别"].isin(NON_PRODUCT)].copy()
svc = c[~(c["_is_prod"] & ~c["物资类别"].isin(NON_PRODUCT))].copy()

# ── 派生字段(对 prod 与 svc 都加) ────────────────────────────────────
ESF = {"M": "市政", "E": "州", "F": "联邦", "N": "国企/全国", "D": "特区"}
tmap = dict(zip(tc.source_text, tc.translated))

def meepp_label(names):
    names = {str(n) for n in names if n and str(n) not in ("None", "nan")}
    excl = [n for n in names if "xclusiv" in n.lower()]
    cota = [n for n in names if "ota" in n.lower() and "reserv" in n.lower()]
    sub = [n for n in names if "ubcontrat" in n.lower()]
    if excl and len(excl) == len(names): return "全部专属ME/EPP(外企不可投)"
    if excl: return "部分专属ME/EPP"
    if cota: return "预留份额"
    if sub: return "要求分包"
    return "无限制"
meepp = (it.groupby("pncp_id")["tipo_beneficio_nome"]
         .apply(lambda s: meepp_label(set(s.dropna()))).rename("ME/EPP限制"))

def portal_link(pncp):
    s = str(pncp or "")
    try:
        left, ano = s.split("/"); parts = left.split("-")
        return f"https://pncp.gov.br/app/editais/{parts[0]}/{ano}/{int(parts[-1])}"
    except Exception:
        return ""
def fix_scheme(u):
    u = str(u or "").strip()
    if not u or u.lower() == "nan": return ""
    return u if u.startswith("http") else "http://" + u

today = date(2026, 6, 18)
def enrich(df):
    df = df.merge(meepp, left_on="pncp_id", right_index=True, how="left")
    df["ME/EPP限制"] = df["ME/EPP限制"].fillna("无限制")
    df["标的中文梗概"] = df["objeto_compra"].map(lambda o: tmap.get(o, ""))
    df["级别"] = df["orgao_esfera_id"].map(ESF).fillna(df["orgao_esfera_id"])
    df["预估金额(万CNY)"] = (df["valor_cny_estimado"] / 1e4).round(1)
    df["预估金额(BRL)"] = df["valor_total_estimado"].round(0)
    df["提交截止日期"] = pd.to_datetime(df["data_encerramento_proposta"]).dt.date
    df["发布日期"] = pd.to_datetime(df["data_publicacao_pncp"]).dt.date
    df["剩余天数"] = df["提交截止日期"].map(lambda d: (d - today).days if pd.notna(d) else None)
    df["时效状态"] = df["剩余天数"].map(lambda d: "未过期" if (d is not None and d >= 0) else "已过期")
    df["价格登记SRP"] = df["srp"].map(lambda x: "是" if x else "否")
    df["长期挂网"] = df["is_long_term_opportunity"].map(lambda x: "是" if x else "否")
    df["PNCP官方链接"] = df["pncp_id"].map(portal_link)
    df["投标链接"] = df["link_sistema_origem"].map(fix_scheme)
    return df
prod, svc = enrich(prod), enrich(svc)

# ── 4) 出表 ───────────────────────────────────────────────────────────
COLS = ["物资类别", "时效状态", "标的中文梗概", "主要采购物品", "预估金额(万CNY)", "预估金额(BRL)",
        "orgao_razao_social", "级别", "uf_sigla", "municipio_nome", "modalidade_nome",
        "ME/EPP限制", "价格登记SRP", "长期挂网", "发布日期", "提交截止日期", "剩余天数",
        "objeto_compra", "投标链接", "PNCP官方链接", "pncp_id"]
REN = {"orgao_razao_social": "采购机构", "uf_sigla": "州", "municipio_nome": "城市",
       "modalidade_nome": "招标方式", "objeto_compra": "标的原文(葡)", "pncp_id": "PNCP编号"}
CORE = ["车辆/工程机械", "工业/生产设备", "IT/电子/通信设备", "医疗器械/耗材", "实验/检测仪器",
        "办公/教学/家具", "厨房/家电", "建材/五金/物料", "燃料/化工品", "药品", "食品/农产品",
        "其他实物", "服务(非实物)", "工程施工"]
def order_export(df):
    df = df.copy()
    df["_o"] = df["物资类别"].map({k: i for i, k in enumerate(CORE)}).fillna(99)
    df = df.sort_values(["_o", "valor_cny_estimado"], ascending=[True, False]).reset_index(drop=True)
    o = df[COLS].rename(columns=REN)
    o.insert(0, "序号", range(1, len(o) + 1))
    return o
prod_x, svc_x = order_export(prod), order_export(svc)

summ = (prod.groupby("物资类别")
        .agg(标数=("pncp_id", "count"),
             未过期标数=("时效状态", lambda s: int((s == "未过期").sum())),
             金额合计_万CNY=("预估金额(万CNY)", "sum"),
             单标中位_万CNY=("预估金额(万CNY)", "median"))
        .reindex([k for k in CORE if k not in NON_PRODUCT]).dropna(how="all").reset_index())
summ["金额合计_万CNY"] = summ["金额合计_万CNY"].round(0)
summ["单标中位_万CNY"] = summ["单标中位_万CNY"].round(1)
summ["标数"] = summ["标数"].astype(int)
summ["未过期标数"] = summ["未过期标数"].astype(int)

with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    prod_x.to_excel(xw, sheet_name="设备产品采购标", index=False)
    summ.to_excel(xw, sheet_name="类别汇总", index=False)
    svc_x.to_excel(xw, sheet_name="剔除_疑似服务工程", index=False)
    wb = xw.book
    widths = {"标的中文梗概": 42, "主要采购物品": 42, "采购机构": 32, "标的原文(葡)": 50,
              "投标链接": 26, "PNCP官方链接": 40, "PNCP编号": 26, "招标方式": 16,
              "ME/EPP限制": 18, "物资类别": 16}
    for sh in wb.worksheets:
        sh.freeze_panes = "A2"
        for col in sh[1]:
            col.font = Font(bold=True, color="FFFFFF")
            col.fill = PatternFill("solid", fgColor="2F5597")
            col.alignment = Alignment(vertical="center", wrap_text=True)
        for j, name in enumerate([cc.value for cc in sh[1]], 1):
            sh.column_dimensions[get_column_letter(j)].width = widths.get(name, 12)

print(f"✅ 设备/产品采购标 {len(prod_x)} 条(候选池 {len(c)},剔疑似服务/工程 {len(svc_x)}) → {OUT}")
print("\n类别分布(产品标):")
print(summ.to_string(index=False))
print(f"\n产品标金额合计:{prod['预估金额(万CNY)'].sum():.0f} 万CNY  |  剩余天数中位 {int(prod['剩余天数'].median())} 天")
print(f"剔除项类别:{svc['物资类别'].value_counts().to_dict()}")
