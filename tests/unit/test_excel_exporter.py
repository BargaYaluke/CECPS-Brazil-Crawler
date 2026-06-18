"""tests/unit/test_excel_exporter.py — Excel 导出单测(lite 版 report.xlsx)。

用内存 DB 造少量数据,导出后用 openpyxl 读回验证 sheet 名、表头、翻译/时效列、
富化列中文化、金额格式。lite 版只导 4 个 sheet:招标主表 / 采购明细 / 过滤审计 / 字段说明。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from src.enrichers.translate import text_hash
from src.storage import (
    Base,
    Contratacao,
    Item,
    TranslationCache,
)


@pytest.fixture
def seeded_engine(monkeypatch: pytest.MonkeyPatch):
    """内存 DB + 少量数据;monkeypatch 让 exporter 的 session_scope 用它。"""
    eng = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(bind=eng)

    with Session(eng) as s:
        s.add(
            Contratacao(
                pncp_id="x-1-001/2026",
                modalidade_nome="Pregão - Eletrônico",
                objeto_compra="Aquisição de medicamentos",
                valor_total_estimado=123456.78,
                valor_cny_estimado=165432.08,
                region_macro="Centro-Oeste",
                region_gdp_tier="middle",
                orgao_razao_social="MINISTÉRIO DA SAÚDE",
                orgao_cnpj="00394544000185",
                uf_sigla="DF",
                data_publicacao_pncp=datetime(2026, 5, 26, 10, 0, 0),
                data_encerramento_proposta=datetime(2099, 12, 31, 0, 0, 0),  # 远期 → 还能投
                is_e_auction=True,
            )
        )
        s.add(
            Contratacao(
                pncp_id="y-1-002/2026",
                modalidade_nome="Concorrência - Eletrônica",
                # 内嵌非法控制字符(换页符 \x0c / 垂直制表符 \x0b)——PNCP 真实数据里出现过
                objeto_compra="Óleo\x0c diesel\x0b LACEN",
                valor_total_estimado=999.0,
                uf_sigla="SP",
                data_encerramento_proposta=datetime(2025, 1, 1, 0, 0, 0),  # 过去 → 已过期
                is_competitive_bid=True,
            )
        )
        s.add(Item(pncp_id="x-1-001/2026", numero_item=1, descricao="Dipirona", valor_total=50.0))
        # 翻译缓存:覆盖 x-1-001 的 objeto(用离线词典绝不会产出的译文,以便区分来源)
        s.add(
            TranslationCache(
                text_hash=text_hash("Aquisição de medicamentos"),
                source_text="Aquisição de medicamentos",
                translated="【缓存】采购药品供医疗使用",
                model="deepseek-v4-flash",
            )
        )
        s.commit()

    # exporter 内部用 session_scope → get_engine();patch 成内存 eng
    monkeypatch.setattr("src.storage.database.get_engine", lambda *a, **kw: eng)
    # filtered_log 目录指到不存在的临时位置(导出空审计 sheet)
    monkeypatch.setattr(
        "src.exporters.excel.get_filter_settings",
        lambda *a, **kw: {"mode": "hard_delete", "filtered_log_dir": "/nonexistent"},
    )
    yield eng
    eng.dispose()


def test_export_creates_lite_sheets(seeded_engine, tmp_path: Path) -> None:
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    stats = export_to_excel(out)

    assert out.exists()
    wb = load_workbook(out)
    assert set(wb.sheetnames) == {"招标主表", "采购明细", "过滤审计", "字段说明"}
    assert stats["招标主表"] == 2
    assert stats["采购明细"] == 1
    assert stats["字段说明"] > 0


def test_illegal_control_chars_are_stripped(seeded_engine, tmp_path: Path) -> None:
    """含控制字符(\\x0c/\\x0b)的 PNCP 文本不会让导出崩,且字符被清掉。"""
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    export_to_excel(out)  # 不抛 IllegalCharacterError 即过第一关
    wb = load_workbook(out)
    ws = wb["招标主表"]
    headers = [c.value for c in ws[1]]
    obj_col = headers.index("标的(原文)") + 1
    texts = [ws.cell(row=r, column=obj_col).value for r in range(2, ws.max_row + 1)]
    target = next((t for t in texts if t and "LACEN" in t), None)
    assert target is not None
    assert "\x0c" not in target and "\x0b" not in target  # 控制字符已剔除
    assert "Óleo diesel LACEN" == target  # 清洗后文本连续


def _col(ws, header: str) -> int:
    return [c.value for c in ws[1]].index(header) + 1


def _col_values(ws, header: str) -> set:
    col = _col(ws, header)
    return {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}


def test_translation_and_deadline(seeded_engine, tmp_path: Path) -> None:
    """标的中文梗概(优先翻译缓存)+ 时效状态 两列正确。"""
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    export_to_excel(out)
    wb = load_workbook(out)
    ws = wb["招标主表"]
    headers = [c.value for c in ws[1]]
    assert "标的(中文梗概)" in headers
    assert "时效状态" in headers

    pncp_col = _col(ws, "PNCP编号")
    row_idx = next(
        r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=pncp_col).value == "x-1-001/2026"
    )
    zh = ws.cell(row=row_idx, column=_col(ws, "标的(中文梗概)")).value
    assert zh == "【缓存】采购药品供医疗使用"  # 优先用翻译缓存(DeepSeek),非离线词典
    assert ws.cell(row=row_idx, column=_col(ws, "时效状态")).value == "无截止"  # 2099 哨兵

    # 时效状态覆盖到"已过期"(y-1-002,2025 截止)
    assert "已过期" in _col_values(ws, "时效状态")


def test_hide_expired_drops_expired_rows(seeded_engine, tmp_path: Path) -> None:
    """--hide-expired:已过期的 y-1-002 被剔除,只剩 x-1-001。"""
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    stats = export_to_excel(out, hide_expired=True)
    assert stats["招标主表"] == 1  # 原 2 条,过期 1 条被隐藏
    wb = load_workbook(out)
    pncp_vals = _col_values(wb["招标主表"], "PNCP编号")
    assert "x-1-001/2026" in pncp_vals
    assert "y-1-002/2026" not in pncp_vals


def test_glossary_sheet_has_terms(seeded_engine, tmp_path: Path) -> None:
    """字段说明 Sheet 解释了关键术语。"""
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    export_to_excel(out)
    wb = load_workbook(out)
    ws = wb["字段说明"]
    fields = _col_values(ws, "字段")
    assert "时效状态" in fields


def test_enrichment_columns_present_and_localized(seeded_engine, tmp_path: Path) -> None:
    """招标主表带富化列(CNY 金额 / 大区 / GDP分层);GDP 分层中文化。"""
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    export_to_excel(out)
    wb = load_workbook(out)
    ws = wb["招标主表"]
    headers = [c.value for c in ws[1]]

    assert "预估金额(CNY)" in headers
    assert "大区" in headers
    assert "GDP分层" in headers

    cny_col = headers.index("预估金额(CNY)") + 1
    assert ws.cell(row=2, column=cny_col).number_format == "#,##0.00"

    gdp_col = headers.index("GDP分层") + 1
    gdp_values = {ws.cell(row=r, column=gdp_col).value for r in range(2, ws.max_row + 1)}
    assert "中" in gdp_values
    assert "middle" not in gdp_values  # 不该出现英文


def test_header_frozen_and_filtered(seeded_engine, tmp_path: Path) -> None:
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    export_to_excel(out)
    wb = load_workbook(out)
    ws = wb["招标主表"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_money_format_applied(seeded_engine, tmp_path: Path) -> None:
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    export_to_excel(out)
    wb = load_workbook(out)
    ws = wb["招标主表"]
    headers = [c.value for c in ws[1]]
    money_col = headers.index("预估金额(BRL)") + 1
    cell = ws.cell(row=2, column=money_col)
    assert cell.number_format == "#,##0.00"


def test_bool_localized(seeded_engine, tmp_path: Path) -> None:
    """布尔列显示 是/否。"""
    from src.exporters.excel import export_to_excel

    out = tmp_path / "report.xlsx"
    export_to_excel(out)
    wb = load_workbook(out)
    ws = wb["招标主表"]
    headers = [c.value for c in ws[1]]
    col = headers.index("电子拍卖") + 1
    values = {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
    assert "是" in values or "否" in values
